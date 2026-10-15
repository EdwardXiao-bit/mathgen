"""M5 里程碑验收（task-42 / T-M5-04）后端判据：FR-T11 列表 / FR-T12 明细 / FR-T13 导出。

判据与任务书的对应（八条验收要求，逐条在这里落地）：

    ① 得分只有一个来源（D80）      → `TestScoreSourceIsSingle` + 列表/明细/CSV/提交响应四点同数
    ② 导出导的是「当前筛选下全部」 → `TestExportAcceptance`（夹具 25 条 > 一页 20）
    ③ CSV 必须是 UTF-8 with BOM    → `TestExportAcceptance`（读回前 3 字节）
    ④ `total == len(items)` 不是自动成立 → 列表↔明细↔库对账 + **真删一行**让它红
    ⑤ 同一时刻翻页不重不漏          → `TestPagingAtTheSameInstant`（把 7 条压成同刻）
    ⑥ 自主练习三字段为 null        → `TestSelfPractice`（+ 界面判据在 UI 模块）
    ⑦ 404 / 422 边界               → `TestAttemptBoundaries`
    ⑧ `q` 的 % 与 _ 按字面量       → `TestAttemptBoundaries`

跑法（在仓库根；`-t backend` 必须在位置参数**之前**）：

    & $VPY -m unittest discover -s backend/tests/e2e_m5 -t backend -v

设计要点见同目录 `_m5_support.py` 的模块文档串：判据是纯函数、世界由夹具造、
负对照把**坏世界**喂给**同一条**判据函数（V9.1/V9.2/V11）。
"""
from __future__ import annotations

import hashlib
import json
import re
import unittest
import uuid
from pathlib import Path

from tests.e2e_m5 import _m5_support as m5

#: 整个模块共用一个真实服务与一份世界（一次 uvicorn 启动，30 秒内跑完）
FIXTURE: dict = {}


def _paper(server: m5.M5Server, *, title: str, count: int) -> dict:
    paper_id = server.new_paper(title=title, count=count)
    published = server.publish(paper_id)
    return {"id": paper_id, "code": published["code"], "title": title, "count": count}


def _submit(server: m5.M5Server, *, student_id: int, paper: dict, correct: int,
            duration_ms: int = 1234) -> dict:
    response = server.submit_graded(
        student_id=student_id, paper_id=paper["id"], uuid_hex=uuid.uuid4().hex[:12],
        correct_count=correct, duration_ms=duration_ms)
    return response.json()


def _build_world(server: m5.M5Server) -> None:
    """造世界。**每一处都顺手自检**：造不出预期世界就当场炸，别让判据建在别的东西上。"""
    papers = {
        "a40": _paper(server, title="M5 成绩验收卷-A（40 题）", count=40),
        "b3": _paper(server, title="M5 成绩验收卷-B（3 题）", count=3),
        "d80": _paper(server, title="M5 成绩验收卷-D（80 题）", count=80),
        "e3": _paper(server, title="M5 成绩验收卷-E（导出覆盖面）", count=3),
    }

    raw_students = [
        ("s1", "肖铭皓", "2024141460455"),
        ("s2", "百分%同学", "2026M5%01"),
        ("s3", "下划_线同学", "2026M5_02"),
        ("s4", "普通甲", "2026M5001"),
    ]
    students = []
    for key, name, student_no in raw_students:
        students.append({"key": key, "student_id": server.student_id(name, student_no),
                         "name": name, "student_no": student_no})
    by_key = {row["key"]: row for row in students}

    attempts = {}
    attempts["a40_23"] = _submit(server, student_id=by_key["s1"]["student_id"],
                                 paper=papers["a40"], correct=23, duration_ms=83000)
    attempts["d80_46"] = _submit(server, student_id=by_key["s1"]["student_id"],
                                 paper=papers["d80"], correct=46, duration_ms=61000)
    attempts["b3_s1_2"] = _submit(server, student_id=by_key["s1"]["student_id"],
                                  paper=papers["b3"], correct=2)
    attempts["b3_s2_1"] = _submit(server, student_id=by_key["s2"]["student_id"],
                                  paper=papers["b3"], correct=1)
    attempts["b3_s3_1"] = _submit(server, student_id=by_key["s3"]["student_id"],
                                  paper=papers["b3"], correct=1)

    # 自主练习（D75）：请求体是平铺的出题参数 + `seed` 查询参数，种子里带回。
    practice = server.request("POST", "/api/student/practice", params={"seed": 777},
                              json_body=dict(m5.CONFIG, count=3), expect=200).json()
    practice_submit = server.request("POST", "/api/attempts", json_body={
        "student_id": by_key["s4"]["student_id"],
        "attempt_uuid": uuid.uuid4().hex[:12],
        "answers": [],
        "duration_ms": 4321,
        "config": dict(m5.CONFIG, count=practice["total"]),
        "seed": practice["seed"],
    }, expect=201).json()
    attempts["practice_s4"] = practice_submit
    FIXTURE["practice_question_count"] = practice["total"]

    # 导出覆盖面夹具：**25 条 > 一页（20）**，否则「导出全部」这条判据在该场景恒真。
    export_students = []
    for index in range(25):
        name, student_no = f"导出学生{index:02d}", f"2026EXP{index:03d}"
        row = {"key": f"e{index}", "student_id": server.student_id(name, student_no),
               "name": name, "student_no": student_no}
        export_students.append(row)
        students.append(row)
        body = _submit(server, student_id=row["student_id"], paper=papers["e3"],
                       correct=index % 4, duration_ms=1000 + index * 7)
        row["attempt_id"] = body["attempt_id"]
        row["expected_score"] = body["score"]
        row["correct"] = body["correct"]
    assert len(export_students) == 25

    FIXTURE.update({
        "papers": papers,
        "students": students,
        "export_students": export_students,
        "attempts": attempts,
    })


def setUpModule():
    server = m5.M5Server()
    server.__enter__()
    FIXTURE["server"] = server
    try:
        _build_world(server)
    except BaseException:
        server.__exit__(None, None, None)
        raise


def tearDownModule():
    server = FIXTURE.pop("server", None)
    if server is not None:
        server.__exit__(None, None, None)


class M5Case(unittest.TestCase):
    """共用世界的基类：`server` / `attempts` / `papers` 都是夹具事实。"""

    @property
    def server(self) -> m5.M5Server:
        return FIXTURE["server"]

    def list_rows(self) -> dict:
        """`{attempt_id: 列表行}`（一次取全，避免每条用例各拉一遍）。"""
        body = self.server.list_attempts(limit=100, offset=0)
        return {int(row["id"]): row for row in body["attempts"]}


class TestContractSurface(unittest.TestCase):
    """判据 0（派单第 1 条）：M5 的契约面**两处副本必须一致**，且 ROUTES 三键真的在。

    `contracts/api.py` 的冻结哈希在仓库里有两份独立誊本：`tests/contract/test_layering.py`
    的 `EXPECTED_SHA256`（那份表里还有 `types.py` / `__init__.py` / `layout.py` / `sheet.py`
    四个别的契约，所以**只能定向取 `api.py` 那一项**）与
    `tests/e2e_m4/test_paper_archive_acceptance.py` 的 `FROZEN_API_SHA`。
    只核对其中一份的话，另一份写成旧值照样全绿 —— 本次验收一开始就撞上了这件事：
    `tests/e2e_m4/REPORT-paper-archive.md` 里还留着 M4 的旧哈希（见 `REPORT-m5.md` §11）。

    这里**不硬编码任何一个哈希**：断言的是「两处副本 == 被冻结文件本身」，
    所以 M6 若合法更新契约（两处一起改 + 文件真变了）这条不会假红。
    """

    BACKEND = Path(__file__).resolve().parents[2]
    FROZEN_FILE = ("app", "contracts", "api.py")
    #: 守卫文件 → 定向正则（组 1 = 该文件里 `api.py` 的那一份副本）
    GUARDS = (
        ("tests/contract/test_layering.py",
         r'"contracts",\s*"api\.py"\)\s*:\s*"([0-9a-f]{64})"'),
        ("tests/e2e_m4/test_paper_archive_acceptance.py",
         r"FROZEN_API_SHA\s*=\s*\"([0-9a-f]{64})\""),
    )

    def test_the_frozen_hash_copies_agree_with_each_other_and_the_file(self):
        actual = hashlib.sha256((self.BACKEND.joinpath(*self.FROZEN_FILE)).read_bytes()).hexdigest()
        copies: dict[str, str] = {}
        for guard, pattern in self.GUARDS:
            source = self.BACKEND.joinpath(*guard.split("/")).read_text(encoding="utf-8")
            match = re.search(pattern, source)
            self.assertIsNotNone(
                match, f"{guard} 里没找到 `api.py` 的冻结哈希 —— 守卫被改写或搬走了")
            copies[guard] = match.group(1)
        self.assertEqual(
            len(set(copies.values())), 1,
            "两处冻结哈希副本互相不一致：" + json.dumps(copies, ensure_ascii=False, indent=2))
        self.assertEqual(
            actual, copies[self.GUARDS[0][0]],
            "两处副本一致，但都不等于当前 `contracts/api.py` 的真实哈希（契约被改而未重新冻结，"
            "或冻结值写错）：两处副本={}".format(copies))

    def test_the_three_m5_routes_exist_and_the_export_route_is_its_own_path(self):
        from app.contracts import api as contract_api

        routes = getattr(contract_api, "ROUTES", None)
        self.assertIsInstance(routes, dict, "契约里没有 ROUTES 表")
        for key in ("list_attempts", "get_attempt", "export_attempts"):
            self.assertIn(key, routes, f"契约 ROUTES 缺 {key}：{sorted(routes)}")
        # 「导出」与「明细」共用 `/api/attempts/...` 前缀：两者路径必须不同，
        # 否则按注册顺序，`/api/attempts/export` 会先被 `{attempt_id}` 吃掉（M5 的真实风险）。
        self.assertNotEqual(
            routes["export_attempts"], routes["get_attempt"],
            f"导出路由与明细路由撞路径：{routes['export_attempts']}")
        self.assertTrue(
            routes["export_attempts"].endswith("/export"),
            f"导出路由不再是独立子路径：{routes['export_attempts']}")
        self.assertIn("{attempt_id}", routes["get_attempt"], routes["get_attempt"])


# ---------------------------------------------------------------------------
# 判据 1（契约侧）：得分只有一个来源
# ---------------------------------------------------------------------------

class TestScoreSourceIsSingle(unittest.TestCase):
    """D80 的契约侧判据。**不 import 产品的算分逻辑**，自己复述公式再对账。"""

    def test_product_score_of_agrees_with_my_independent_restatement(self):
        from app.contracts.api import score_of

        mismatches = [
            (total, correct, score_of(total, correct), m5.score_by_contract(total, correct))
            for total in range(0, 201)
            for correct in range(-2, total + 3)
            if score_of(total, correct) != m5.score_by_contract(total, correct)
        ]
        self.assertEqual([], mismatches,
                         f"产品 score_of 与我的独立复述在 {len(mismatches)} 个点上不一致：{mismatches[:5]}")

    def test_divergence_points_are_exactly_the_documented_ones(self):
        """文档说「8 个点、真实题量内 2 个」——这句话我也要自己算一遍。"""
        all_points = m5.divergence_points(200)
        self.assertEqual(8, len(all_points),
                         f"total<=200 的分叉点数不是 8：{[(p['total'], p['correct']) for p in all_points]}")
        real_points = m5.divergence_points(200, limit=100)
        self.assertEqual([(40, 23), (80, 46)], [(p["total"], p["correct"]) for p in real_points])
        for point in all_points:
            self.assertTrue(point["exact_half"],
                            f"{point['total']}/{point['correct']} 不是「精确值正好 .5」：{point['float']!r}")
            self.assertLess(point["float"], point["contract"],
                            "分叉点的浮点值应当比精确值略低（这正是旧算法少 1 分的原因）")

    def test_the_two_real_world_points_are_where_a_local_algorithm_would_lie(self):
        """真实题量内的两个分叉点：界面若自己算就会显示 57，而后端给 58。"""
        for total, correct in ((40, 23), (80, 46)):
            with self.subTest(total=total, correct=correct):
                self.assertEqual(57, m5.score_by_old_js(total, correct))
                self.assertEqual(58, m5.score_by_contract(total, correct))
                self.assertEqual([], m5.score_source_problems(
                    total=total, correct=correct, observed={"后端 score": 58}))

    def test_negative_control_a_locally_computed_score_is_caught(self):
        problems = m5.score_source_problems(total=40, correct=23, observed={"页面显示的分数": 57})
        self.assertTrue(problems, "界面自己算出来的 57 必须被这条判据抓住")
        joined = " | ".join(problems)
        self.assertIn("58", joined)
        self.assertIn("D80", joined)

    def test_negative_control_no_observation_is_caught(self):
        problems = m5.score_source_problems(total=40, correct=23, observed={})
        self.assertTrue(problems, "没有任何观测点时判据必须报问题（空集对空集最容易假绿）")

    def test_negative_control_missing_inputs_are_caught(self):
        problems = m5.score_source_problems(total=None, correct=None, observed={"列表": 58})
        self.assertTrue(problems)
        self.assertIn("无法核对", " | ".join(problems))


class TestAnInjectedDeviationIsCaught(unittest.TestCase):
    """**注入式**负对照（board 卡第 6 条）：让产品代码自己下发一个偏离 `score_of()` 的 `score`。

    手法：进程内 app（只有同进程才能 patch）+ `unittest.mock.patch.object(attempt_service,
    "score_of", …)` 临时把服务层的算分函数换成「永远给 999 分」。同一个世界、同一条判据
    （`m5.score_source_problems`）在**注入前**与**注入后**各跑一次，两态原始输出都落进
    `artifacts/task42-injected-deviation.json` —— 证明这条判据不是只能对空集假绿。

    为什么不用真机 uvicorn：真实服务是**另一个进程**，`mock.patch` 打不进去（这正是
    §7 方法边界里写的事）。所以这一条走 TestClient，且只用来证明「判据会红」。
    """

    TOTAL, CORRECT = 4, 2          # 契约口径：score_of(4, 2) = 50
    LIE = 999                      # 任何口径都算不出来的分数

    def _world(self, app: m5.InProcessApp) -> tuple[int, int, dict]:
        from sqlalchemy import text as sql_text

        paper_id = app.make_paper(count=self.TOTAL)
        app.publish(paper_id)
        created = app.client.post("/api/students",
                                  json={"name": "注入同学", "student_no": "2026M5INJ"})
        self.assertIn(created.status_code, (200, 201), created.text)
        student_id = int(created.json()["student_id"])
        with app.session() as session:
            rows = session.execute(
                sql_text("SELECT seq, payload FROM paper_item WHERE paper_id=:pid ORDER BY seq"),
                {"pid": paper_id}).all()
        questions = {int(row[0]): json.loads(row[1])["answer"] for row in rows}
        self.assertEqual(self.TOTAL, len(questions), questions)
        return paper_id, student_id, questions

    def _observe(self, app: m5.InProcessApp, *, paper_id: int, student_id: int,
                 questions: dict, tag: str) -> dict:
        """同一批读法：提交响应 / 列表 / 明细三处的 `score`。"""
        created = app.client.post("/api/attempts", json={
            "student_id": student_id,
            "attempt_uuid": f"injected-deviation-{tag}",
            "answers": m5.answers_with(questions, correct_count=self.CORRECT),
            "duration_ms": 1234,
            "paper_id": paper_id,
        })
        self.assertIn(created.status_code, (200, 201), created.text)
        body = created.json()
        attempt_id = int(body["attempt_id"])
        listing = app.client.get("/api/attempts",
                                 params={"limit": 20, "offset": 0, "paper_id": paper_id}).json()
        row = next((item for item in listing["attempts"] if int(item["id"]) == attempt_id), None)
        self.assertIsNotNone(row, f"列表里没有刚提交的 attempt {attempt_id}：{listing}")
        detail = app.client.get(f"/api/attempts/{attempt_id}").json()
        self.assertEqual((self.TOTAL, self.CORRECT), (body.get("total"), body.get("correct")),
                         "注入的是分数，不该动摇 total/correct")
        return {"提交响应": body["score"], "列表": row["score"], "明细": detail["score"]}

    def test_a_backend_that_lies_about_score_is_named_by_the_same_checker(self):
        from unittest import mock

        from app.services import attempt_service

        with m5.InProcessApp() as app:
            paper_id, student_id, questions = self._world(app)

            clean = self._observe(app, paper_id=paper_id, student_id=student_id,
                                  questions=questions, tag="clean")
            clean_problems = m5.score_source_problems(
                total=self.TOTAL, correct=self.CORRECT, observed=clean)

            with mock.patch.object(attempt_service, "score_of", lambda total, correct: self.LIE):
                lying = self._observe(app, paper_id=paper_id, student_id=student_id,
                                      questions=questions, tag="lying")
                lying_problems = m5.score_source_problems(
                    total=self.TOTAL, correct=self.CORRECT, observed=lying)

            m5.write_evidence("task42-injected-deviation.json", {
                "patched": "app.services.attempt_service.score_of -> 999（永远多给分）",
                "total": self.TOTAL, "correct": self.CORRECT,
                "contract_value": m5.score_by_contract(self.TOTAL, self.CORRECT),
                "before_injection": {"observed": clean, "problems": clean_problems},
                "after_injection": {"observed": lying, "problems": lying_problems},
            })

            self.assertEqual({}, {k: v for k, v in clean.items() if v != 50},
                             f"注入前世界就不干净：{clean}")
            self.assertEqual([], clean_problems, clean_problems)
            self.assertEqual(3, len(lying_problems),
                             f"注入后三处观测都该被点名，实际只报了：{lying_problems}")
            joined = " | ".join(lying_problems)
            for label in ("提交响应", "列表", "明细"):
                self.assertIn(label, joined, joined)
            self.assertIn("999", joined, joined)


# ---------------------------------------------------------------------------
# 判据 2/4：列表 ↔ 明细 ↔ 库
# ---------------------------------------------------------------------------

class TestListDetailReconciliation(M5Case):

    def test_every_attempt_reconciles_with_detail_and_database(self):
        db_path = self.server.db_path
        rows = self.list_rows()
        self.assertGreaterEqual(len(rows), 30, "夹具里的成绩条数异常，先查夹具")
        observations, problems = [], []
        for attempt_id, row in sorted(rows.items()):
            detail = self.server.get_attempt(attempt_id)
            db_row = m5.db_attempt_row(db_path, attempt_id)
            item_count = m5.readonly_count(db_path, "attempt_item", "attempt_id=?", (attempt_id,))
            paper_row = (m5.db_paper_row(db_path, int(row["paper_id"]))
                         if row.get("paper_id") is not None else None)
            found = (m5.list_detail_problems(summary=row, detail=detail)
                     + m5.db_cross_check_problems(summary=row, row=db_row,
                                                  item_count=item_count, paper_row=paper_row))
            observations.append({
                "attempt_id": attempt_id,
                "list": {key: row.get(key) for key in
                         ("total", "correct", "score", "paper_id", "paper_code", "submitted_at")},
                "detail_items": len(detail.get("items") or []),
                "detail_total": detail.get("total"),
                "db_items": item_count,
                "problems": found,
            })
            problems.extend(f"attempt {attempt_id}: {item}" for item in found)
        m5.write_evidence("task42-list-detail-reconciliation.json", {
            "attempts": len(rows),
            "problems": problems,
            "observations": observations,
        })
        self.assertEqual([], problems, "列表/明细/库三方对账出现分歧")

    def test_score_is_one_number_across_submit_list_detail_and_csv(self):
        server = self.server
        target = FIXTURE["attempts"]["a40_23"]
        paper_id = FIXTURE["papers"]["a40"]["id"]
        row = self.list_rows()[int(target["attempt_id"])]
        detail = server.get_attempt(target["attempt_id"])
        csv_bytes, _ = server.export_bytes(format="csv", paper_id=paper_id)
        _, csv_rows = m5.csv_records(csv_bytes)
        csv_row = m5.csv_row_for_student(csv_rows, "2024141460455")
        self.assertIsNotNone(csv_row, "CSV 里找不到这位学生的成绩行")

        observed = {
            "提交响应 score": target["score"],
            "列表 score": row["score"],
            "明细 score": detail["score"],
            "CSV 得分列": int(csv_row[4]),
        }
        problems = m5.score_source_problems(total=row["total"], correct=row["correct"],
                                            observed=observed)
        m5.write_evidence("task42-score-one-source.json", {
            "total": row["total"], "correct": row["correct"], "observed": observed,
            "old_js_would_say": m5.score_by_old_js(row["total"], row["correct"]),
            "problems": problems,
        })
        self.assertEqual([], problems)
        self.assertEqual(58, target["score"], "40 题答对 23 题按 D80 必须是 58")

    def test_the_second_real_world_point_is_also_58_on_the_real_backend(self):
        target = FIXTURE["attempts"]["d80_46"]
        row = self.list_rows()[int(target["attempt_id"])]
        problems = m5.score_source_problems(total=row["total"], correct=row["correct"],
                                            observed={"列表 score": row["score"]})
        self.assertEqual([], problems)
        self.assertEqual(58, row["score"])
        self.assertEqual(57, m5.score_by_old_js(80, 46))

    # -- 负对照：把坏世界喂给同一条判据 ---------------------------------

    def test_negative_control_a_swapped_correct_count_is_caught(self):
        target = FIXTURE["attempts"]["a40_23"]
        row = dict(self.list_rows()[int(target["attempt_id"])])
        detail = self.server.get_attempt(target["attempt_id"])
        row["correct"] = row["correct"] - 1
        self.assertNotEqual(row["correct"], detail["correct"], "世界没被改坏，负对照无效")
        problems = m5.list_detail_problems(summary=row, detail=detail)
        self.assertTrue(problems, "改小 correct 后同一条判据必须变红")
        joined = " | ".join(problems)
        self.assertIn(str(row["correct"]), joined)
        self.assertIn(str(detail["correct"]), joined)

    def test_negative_control_a_truncated_detail_is_caught(self):
        """`total == len(items)` 不是自动成立：明细少一题必须当场红。"""
        target = FIXTURE["attempts"]["a40_23"]
        row = self.list_rows()[int(target["attempt_id"])]
        detail = dict(self.server.get_attempt(target["attempt_id"]))
        detail["items"] = detail["items"][:-1]
        problems = m5.list_detail_problems(summary=row, detail=detail)
        self.assertTrue(problems)
        joined = " | ".join(problems)
        self.assertIn(str(len(detail["items"])), joined)
        self.assertIn(str(detail["total"]), joined)

    def test_negative_control_a_locally_computed_row_score_is_caught(self):
        target = FIXTURE["attempts"]["a40_23"]
        row = dict(self.list_rows()[int(target["attempt_id"])])
        detail = self.server.get_attempt(target["attempt_id"])
        row["score"] = m5.score_by_old_js(row["total"], row["correct"])
        problems = m5.list_detail_problems(summary=row, detail=detail)
        self.assertTrue(problems, "列表行自己算出的 57 必须被抓住")
        self.assertIn("D80", " | ".join(problems))

    def test_deleting_a_real_attempt_item_makes_the_same_checker_red(self):
        """判据 4 的硬证据：**真删**库里一行 `attempt_item`，同一条判据必须红并点名两个数。

        删行会污染同进程共享的世界（`test_every_attempt_reconciles…` 按字母序排在后面，
        且遍历**所有**成绩），所以这里必须 `finally` 把那一行原样插回，并复验它重新变绿。
        """
        server = self.server
        paper = _paper(server, title="M5 成绩验收卷-F（删行实验）", count=5)
        student_id = server.student_id("删行实验同学", "2026M5DEL01")
        body = _submit(server, student_id=student_id, paper=paper, correct=3)
        attempt_id = int(body["attempt_id"])
        row = self.list_rows()[attempt_id]
        before = server.get_attempt(attempt_id)
        self.assertEqual([], m5.list_detail_problems(summary=row, detail=before),
                         "删之前这条判据本来就该是绿的")

        victim = m5.readonly_rows(
            server.db_path,
            "SELECT id, attempt_id, seq, payload, student_answer, is_correct"
            " FROM attempt_item WHERE attempt_id=? ORDER BY id LIMIT 1",
            (attempt_id,))[0]
        try:
            deleted = m5.direct_write(
                server.db_path,
                "DELETE FROM attempt_item WHERE id=?",
                (victim["id"],))
            self.assertEqual(1, deleted, "没真的删掉一行，负对照无效")

            after = server.get_attempt(attempt_id)
            problems = m5.list_detail_problems(summary=row, detail=after)
            item_count = m5.readonly_count(server.db_path, "attempt_item",
                                           "attempt_id=?", (attempt_id,))
            db_problems = m5.db_cross_check_problems(
                summary=row, row=m5.db_attempt_row(server.db_path, attempt_id),
                item_count=item_count)
        finally:
            m5.direct_write(
                server.db_path,
                "INSERT INTO attempt_item"
                " (id, attempt_id, seq, payload, student_answer, is_correct)"
                " VALUES (?, ?, ?, ?, ?, ?)",
                (victim["id"], victim["attempt_id"], victim["seq"], victim["payload"],
                 victim["student_answer"], victim["is_correct"]))

        m5.write_evidence("task42-real-deletion-negative-control.json", {
            "attempt_id": attempt_id,
            "items_before": len(before["items"]), "items_after": len(after["items"]),
            "detail_total": after["total"], "db_item_rows": item_count,
            "problems": problems, "db_problems": db_problems,
        })
        self.assertTrue(problems, "真删一行 attempt_item 后同一条判据必须变红")
        joined = " | ".join(problems)
        self.assertIn(str(len(after["items"])), joined)
        self.assertIn(str(after["total"]), joined)
        self.assertTrue(db_problems, "库侧对账也必须红（item 行数与 attempt.total 不一致）")

        restored = server.get_attempt(attempt_id)
        self.assertEqual([], m5.list_detail_problems(summary=row, detail=restored),
                         "删行用例没把世界恢复原样，后面的全量对账会被污染")


# ---------------------------------------------------------------------------
# 判据 5：同一时刻的多条记录，翻页不重不漏
# ---------------------------------------------------------------------------

class TestPagingAtTheSameInstant(unittest.TestCase):
    """把 `submitted_at` 压成**完全相同**，才真正在测「同一时刻按 id 倒序」的稳定分页。

    真实提交不可能产生同一微秒的多条记录，所以这里走进程内（可改库）——这是**构造**
    判据需要的世界，不是往夹具里灌数据凑覆盖面。
    """

    SAME_INSTANT = "2026-10-15T10:00:00+08:00"
    COUNT = 7

    def _world(self, app: m5.InProcessApp) -> tuple[int, list[int]]:
        from sqlalchemy import update

        from app.models import Attempt

        paper_id = app.make_paper(count=3)
        app.publish(paper_id)
        created = app.client.post("/api/students",
                                  json={"name": "同刻同学", "student_no": "2026M5SAME"})
        self.assertIn(created.status_code, (200, 201), created.text)
        student_id = int(created.json()["student_id"])

        attempt_ids = []
        for index in range(self.COUNT):
            response = app.client.post("/api/attempts", json={
                "student_id": student_id,
                "attempt_uuid": f"same-instant-{index}",
                "answers": [],
                "duration_ms": 1000 + index,
                "paper_id": paper_id,
            })
            self.assertIn(response.status_code, (200, 201), response.text)
            attempt_ids.append(int(response.json()["attempt_id"]))
        self.assertEqual(self.COUNT, len(set(attempt_ids)))

        with app.session() as session:
            session.execute(update(Attempt).values(submitted_at=self.SAME_INSTANT))
            session.commit()
        return paper_id, attempt_ids

    def test_same_instant_rows_paginate_without_gaps_or_duplicates(self):
        with m5.InProcessApp() as app:
            paper_id, attempt_ids = self._world(app)
            pages = [app.client.get("/api/attempts",
                                    params={"limit": 2, "offset": offset, "paper_id": paper_id}).json()
                     for offset in (0, 2, 4, 6)]
            # 契约排序：`submitted_at desc, id desc`。同刻时全部由 id 倒序决定 —— 这里
            # **独立复述**一遍，而不是从接口拿顺序再拿它当期望（那就是同义反复）。
            expected_ids = sorted(attempt_ids, reverse=True)
            problems = m5.paging_problems(pages=pages, expected_ids=expected_ids,
                                          expected_total=self.COUNT)
            observed_instants = {page["attempts"][0]["submitted_at"] for page in pages if page["attempts"]}
            m5.write_evidence("task42-same-instant-paging.json", {
                "same_instant": self.SAME_INSTANT,
                "attempt_ids": attempt_ids,
                "expected_order": expected_ids,
                "page_ids": [[row["id"] for row in page["attempts"]] for page in pages],
                "observed_submitted_at": sorted(observed_instants),
                "problems": problems,
            })
            self.assertEqual({self.SAME_INSTANT}, observed_instants,
                             "夹具没把 7 条压成同一时刻，这条判据就退化成普通翻页")
            self.assertEqual([], problems)

    def test_negative_control_a_dropped_page_is_caught(self):
        pages = [
            {"total": 7, "attempts": [{"id": 7}, {"id": 6}]},
            {"total": 7, "attempts": [{"id": 5}, {"id": 4}]},
            {"total": 7, "attempts": [{"id": 1}]},
        ]
        problems = m5.paging_problems(pages=pages, expected_ids=[7, 6, 5, 4, 3, 2, 1], expected_total=7)
        self.assertTrue(problems, "少了一页必须红")
        self.assertIn("缺", " | ".join(problems))

    def test_negative_control_a_duplicated_row_is_caught(self):
        pages = [
            {"total": 7, "attempts": [{"id": 7}, {"id": 6}]},
            {"total": 7, "attempts": [{"id": 6}, {"id": 5}]},
        ]
        problems = m5.paging_problems(pages=pages, expected_ids=[7, 6, 5], expected_total=7)
        self.assertTrue(problems, "重复行必须红")
        self.assertIn("重复", " | ".join(problems))

    def test_negative_control_a_page_that_reports_a_different_total_is_caught(self):
        pages = [{"total": 7, "attempts": [{"id": 7}]}, {"total": 6, "attempts": [{"id": 6}]}]
        problems = m5.paging_problems(pages=pages, expected_ids=[7, 6], expected_total=7)
        self.assertTrue(problems)
        self.assertIn("total", " | ".join(problems))


# ---------------------------------------------------------------------------
# 判据 7/8：404 / 422 边界与 `q` 的字面量语义
# ---------------------------------------------------------------------------

class TestAttemptBoundaries(M5Case):

    def test_missing_attempt_is_404_with_the_contract_code(self):
        response = self.server.get_attempt_response(999999)
        self.assertEqual(404, response.status_code, response.text)
        body = response.json()
        # 契约 §5.5 / docs/04-接口契约.md:300 的错误信封：{"error": {"code", "message"}}
        error = body.get("error") or {}
        self.assertEqual("ATTEMPT_NOT_FOUND", error.get("code"), body)
        self.assertIn("不存在", error.get("message", ""))
        m5.write_evidence("task42-boundary-404.json", {
            "status": response.status_code, "body": body,
        })

    def test_limit_and_offset_out_of_range_are_422(self):
        observed = {}
        for bad in (0, -1, 101, 1000):
            response = self.server.list_attempts_response(limit=bad, expect=None)
            observed[f"limit={bad}"] = response.status_code
            self.assertEqual(422, response.status_code, f"limit={bad} 应当 422：{response.text[:200]}")
        for good in (1, 100):
            response = self.server.list_attempts_response(limit=good)
            observed[f"limit={good}"] = response.status_code
            self.assertEqual(200, response.status_code, f"limit={good} 应当 200")
        negative_offset = self.server.list_attempts_response(offset=-1, expect=None)
        observed["offset=-1"] = negative_offset.status_code
        self.assertEqual(422, negative_offset.status_code)
        m5.write_evidence("task42-boundary-limits.json", observed)

    def test_limit_out_of_range_is_rejected_by_the_schema_not_by_luck(self):
        """422 必须是**参数校验**给的（有 error 信封），不是服务端 500 顺带出来的。"""
        response = self.server.list_attempts_response(limit=0, expect=None)
        self.assertEqual(422, response.status_code)
        body = response.json()
        error = body.get("error") or {}
        self.assertEqual("VALIDATION_ERROR", error.get("code"), body)
        self.assertIn("limit", error.get("message", ""), body)

    def test_q_wildcards_are_literal(self):
        students = FIXTURE["students"]
        with_percent = [row for row in students if "%" in row["name"] or "%" in row["student_no"]]
        with_underscore = [row for row in students if "_" in row["name"] or "_" in row["student_no"]]
        self.assertTrue(with_percent, "夹具里没有含 % 的学生，这条判据没在守护任何东西")
        self.assertTrue(with_underscore, "夹具里没有含 _ 的学生，这条判据没在守护任何东西")

        observed = {}
        for char in ("%", "_"):
            body = self.server.list_attempts(q=char, limit=100, offset=0)
            returned = [int(row["student_id"]) for row in body["attempts"]]
            observed[f"q={char}"] = {
                "returned_student_ids": sorted(set(returned)),
                "returned_rows": len(body["attempts"]),
                "total_all": self.server.list_attempts(limit=100, offset=0)["total"],
            }
            problems = m5.wildcard_problems(char=char, students=students, returned_ids=returned)
            self.assertEqual([], problems, f"q={char} 的字面量语义不对")
            self.assertTrue(returned, f"q={char} 返回了空集 —— 判据无法区分「对」与「全都不过滤」")

        # 空串是合法的「不筛选」（契约没给 min_length）。
        empty = self.server.list_attempts(q="", limit=100, offset=0)
        observed["q="] = {"total": empty["total"]}
        self.assertEqual(self.server.list_attempts(limit=100, offset=0)["total"], empty["total"])
        m5.write_evidence("task42-q-wildcards.json", observed)

    def test_negative_control_an_unescaped_wildcard_world_is_caught(self):
        students = FIXTURE["students"]
        returned = [row["student_id"] for row in students]      # 「忘了转义 → 匹配全部」
        problems = m5.wildcard_problems(char="%", students=students, returned_ids=returned)
        self.assertTrue(problems, "把 % 当通配符的世界必须被抓住")
        self.assertIn("通配符", " | ".join(problems))

    def test_paper_filter_is_exactly_that_paper(self):
        paper = FIXTURE["papers"]["e3"]
        body = self.server.list_attempts(paper_id=paper["id"], limit=100, offset=0)
        self.assertEqual(25, body["total"])
        self.assertEqual({paper["id"]}, {row["paper_id"] for row in body["attempts"]})


# ---------------------------------------------------------------------------
# 判据 6：自主练习（D75）
# ---------------------------------------------------------------------------

class TestSelfPractice(M5Case):

    def test_practice_attempt_has_null_paper_fields_and_no_display_wording(self):
        attempt_id = int(FIXTURE["attempts"]["practice_s4"]["attempt_id"])
        row = self.list_rows()[attempt_id]
        detail = self.server.get_attempt(attempt_id)
        raw_list = self.server.list_attempts_response(limit=100, offset=0).text
        raw_detail = self.server.get_attempt_response(attempt_id).text

        problems = (m5.practice_contract_problems(summary=row, payload_text=raw_list)
                    + m5.practice_contract_problems(summary=detail, payload_text=raw_detail)
                    + m5.list_detail_problems(summary=row, detail=detail)
                    + m5.db_cross_check_problems(
                        summary=row, row=m5.db_attempt_row(self.server.db_path, attempt_id),
                        item_count=m5.readonly_count(self.server.db_path, "attempt_item",
                                                     "attempt_id=?", (attempt_id,))))
        m5.write_evidence("task42-self-practice.json", {
            "attempt_id": attempt_id,
            "list_row": {key: row.get(key) for key in
                         ("paper_id", "paper_code", "paper_title", "total", "correct", "score")},
            "detail_paper_fields": {key: detail.get(key) for key in
                                    ("paper_id", "paper_code", "paper_title")},
            "list_body_has_display_wording": "自主练习" in raw_list,
            "problems": problems,
        })
        self.assertEqual([], problems)
        self.assertTrue(row["paper_id"] is None and row["paper_code"] is None
                        and row["paper_title"] is None,
                        f"自主练习的卷字段必须是 null：{row}")

    def test_negative_control_a_practice_label_from_the_backend_is_caught(self):
        attempt_id = int(FIXTURE["attempts"]["practice_s4"]["attempt_id"])
        row = dict(self.list_rows()[attempt_id])
        row["paper_title"] = "自主练习"
        problems = m5.practice_contract_problems(summary=row,
                                                 payload_text=json.dumps(row, ensure_ascii=False))
        self.assertTrue(problems, "后端返回显示文案必须被抓住")
        self.assertIn("自主练习", " | ".join(problems))

    def test_practice_scores_through_the_same_single_source(self):
        attempt_id = int(FIXTURE["attempts"]["practice_s4"]["attempt_id"])
        row = self.list_rows()[attempt_id]
        detail = self.server.get_attempt(attempt_id)
        problems = m5.score_source_problems(total=row["total"], correct=row["correct"],
                                            observed={"列表": row["score"], "明细": detail["score"]})
        self.assertEqual([], problems)


# ---------------------------------------------------------------------------
# 判据 2/3：导出（当前筛选下全部 + CSV BOM）
# ---------------------------------------------------------------------------

class TestExportAcceptance(M5Case):

    def _csv_for_paper(self, paper_id: int) -> tuple[bytes, list[list[str]], dict]:
        listing = self.server.list_attempts(paper_id=paper_id, limit=m5.PAGE_LIMIT, offset=0)
        data, headers = self.server.export_bytes(format="csv", paper_id=paper_id)
        problems, rows = m5.csv_problems(data=data, expected_rows=listing["total"])
        return data, rows, {"listing": listing, "headers": headers, "problems": problems}

    def test_export_covers_everything_under_the_filter(self):
        paper_id = FIXTURE["papers"]["e3"]["id"]
        data, rows, extra = self._csv_for_paper(paper_id)
        listing = extra["listing"]
        # 断言用的是**响应里的** limit/total，不写死 20：限流口径由服务端回显。
        coverage = m5.export_coverage_problems(csv_rows=len(rows),
                                               list_total=listing["total"],
                                               limit=listing["limit"])
        m5.write_evidence("task42-export-coverage.json", {
            "paper_id": paper_id,
            "list_total": listing["total"],
            "list_limit_echoed": listing["limit"],
            "list_rows_on_page": len(listing["attempts"]),
            "csv_rows": len(rows),
            "csv_problems": extra["problems"],
            "coverage_problems": coverage,
            "content_disposition": extra["headers"].get("content-disposition"),
            "content_type": extra["headers"].get("content-type"),
        })
        self.assertEqual(25, listing["total"], "夹具里的导出覆盖面卷应当有 25 条成绩")
        self.assertEqual(20, listing["limit"])
        self.assertEqual(20, len(listing["attempts"]), "列表一页只给 20 条")
        self.assertEqual([], extra["problems"])
        self.assertEqual([], coverage)

    def test_negative_control_an_export_that_respected_the_page_is_caught(self):
        problems = m5.export_coverage_problems(csv_rows=20, list_total=25, limit=20)
        self.assertTrue(problems, "只导一页的导出必须被抓住")
        self.assertIn("分页", " | ".join(problems))

    def test_negative_control_a_fixture_smaller_than_one_page_is_caught(self):
        """夹具自己比一页还小的时候，「导出全部」这条判据在那个场景里恒真 —— 必须报出来。"""
        problems = m5.export_coverage_problems(csv_rows=3, list_total=3, limit=20)
        self.assertTrue(problems)
        self.assertIn("夹具失效", " | ".join(problems))

    def test_csv_starts_with_a_utf8_bom(self):
        data, _rows, _extra = self._csv_for_paper(FIXTURE["papers"]["e3"]["id"])
        self.assertEqual(m5.BOM, data[:3], f"CSV 前 3 字节是 {data[:3]!r}")
        self.assertTrue(data.decode("utf-8-sig").startswith("卷号,卷名,姓名,学号,得分"),
                        data.decode("utf-8-sig")[:80])

    def test_negative_control_a_csv_without_the_bom_is_caught(self):
        data, _rows, _extra = self._csv_for_paper(FIXTURE["papers"]["e3"]["id"])
        stripped = data[3:]
        problems, rows = m5.csv_problems(data=stripped, expected_rows=25)
        self.assertTrue(problems, "去掉 BOM 必须红")
        self.assertIn("BOM", " | ".join(problems))
        # 反方向：同一份内容加回 BOM 就是干净的 —— 证明这条判据不是恒红。
        clean_problems, clean_rows = m5.csv_problems(data=data, expected_rows=25)
        self.assertEqual([], clean_problems)
        self.assertEqual(len(rows), len(clean_rows))

    def test_negative_control_a_wrong_header_order_is_caught(self):
        data, rows, _extra = self._csv_for_paper(FIXTURE["papers"]["e3"]["id"])
        swapped = data.replace("得分,答对题数".encode("utf-8"), "答对题数,得分".encode("utf-8"))
        self.assertNotEqual(data, swapped, "世界没被改坏，负对照无效")
        problems, _ = m5.csv_problems(data=swapped, expected_rows=25)
        self.assertTrue(problems)
        self.assertIn("表头", " | ".join(problems))

    def test_negative_control_a_truncated_csv_is_caught(self):
        data, _rows, _extra = self._csv_for_paper(FIXTURE["papers"]["e3"]["id"])
        keep = data.decode("utf-8-sig").splitlines()
        truncated = m5.BOM + ("\n".join(keep[:11]) + "\n").encode("utf-8")
        problems, _ = m5.csv_problems(data=truncated, expected_rows=25)
        self.assertTrue(problems)
        self.assertIn("数据行数", " | ".join(problems))

    def test_export_rows_carry_the_contract_numbers(self):
        paper = FIXTURE["papers"]["e3"]
        _data, rows, _extra = self._csv_for_paper(paper["id"])
        by_no = {row[3]: row for row in rows}
        problems = []
        for student in FIXTURE["export_students"]:
            row = by_no.get(student["student_no"])
            if row is None:
                problems.append(f"CSV 缺学生 {student['student_no']}")
                continue
            if row[0] != paper["code"]:
                problems.append(f"{student['student_no']} 卷号列={row[0]!r}，应为 {paper['code']!r}")
            if row[1] != paper["title"]:
                problems.append(f"{student['student_no']} 卷名列={row[1]!r}")
            if row[2] != student["name"]:
                problems.append(f"{student['student_no']} 姓名列={row[2]!r}")
            if row[5] != str(student["correct"]):
                problems.append(f"{student['student_no']} 答对题数={row[5]!r}，应为 {student['correct']}")
            if row[6] != "3":
                problems.append(f"{student['student_no']} 总题数={row[6]!r}，应为 3")
            expected = m5.score_by_contract(3, student["correct"])
            if int(row[4]) != expected:
                problems.append(f"{student['student_no']} 得分={row[4]!r}，按 D80 应为 {expected}")
            if not row[7] or not row[7].endswith("秒"):
                problems.append(f"{student['student_no']} 用时列={row[7]!r}，应是人可读形式（…秒 / …分…秒）")
            if not row[8]:
                problems.append(f"{student['student_no']} 提交时间列是空的")
        m5.write_evidence("task42-export-rows.json", {"rows": len(rows), "problems": problems})
        self.assertEqual([], problems)

    def test_export_filters_match_the_list_filters(self):
        paper = FIXTURE["papers"]["e3"]
        student = FIXTURE["export_students"][7]
        listing = self.server.list_attempts(q=student["student_no"], limit=100, offset=0)
        self.assertEqual(1, listing["total"])
        data, _headers = self.server.export_bytes(format="csv", q=student["student_no"])
        problems, rows = m5.csv_problems(data=data, expected_rows=1)
        self.assertEqual([], problems)
        self.assertEqual(student["student_no"], rows[0][3])
        # 按卷筛选时，导出里不能混进别的卷 —— 逐行核对卷号列。
        data2, _ = self.server.export_bytes(format="csv", paper_id=paper["id"])
        problems2, rows2 = m5.csv_problems(data=data2, expected_rows=25)
        self.assertEqual([], problems2)
        self.assertEqual({paper["code"]}, {row[0] for row in rows2})
        self.assertIn(student["student_no"], {row[3] for row in rows2})

    def test_xlsx_export_is_a_real_workbook(self):
        paper_id = FIXTURE["papers"]["e3"]["id"]
        data, headers = self.server.export_bytes(format="xlsx", paper_id=paper_id)
        self.assertEqual(b"PK\x03\x04", data[:4], f"xlsx 不是 zip 容器：{data[:8]!r}")
        self.assertIn("spreadsheetml", headers.get("content-type", ""))
        import io

        from openpyxl import load_workbook

        workbook = load_workbook(io.BytesIO(data))
        sheet = workbook.active
        header = [cell.value for cell in sheet[1]]
        self.assertEqual(list(m5.COLUMN_LABELS), header)
        self.assertEqual(26, sheet.max_row, "表头 + 25 条数据")
        self.assertEqual(9, sheet.max_column)
        m5.write_evidence("task42-export-xlsx.json", {
            "bytes": len(data), "header": header,
            "max_row": sheet.max_row, "max_column": sheet.max_column,
            "freeze_panes": sheet.freeze_panes,
            "content_disposition": headers.get("content-disposition"),
            "content_type": headers.get("content-type"),
        })

    def test_invalid_format_is_422(self):
        response = self.server.export_response(format="pdf", expect=None)
        self.assertEqual(422, response.status_code, response.text)
        body = response.json()
        error = body.get("error") or {}
        self.assertEqual("VALIDATION_ERROR", error.get("code"), body)
        self.assertIn("pdf", error.get("message", ""), body)
        m5.write_evidence("task42-export-invalid-format.json",
                          {"status": response.status_code, "body": body})

    def test_export_filename_is_the_contract_shape(self):
        data, headers = self.server.export_bytes(format="csv", paper_id=FIXTURE["papers"]["e3"]["id"])
        disposition = headers.get("content-disposition", "")
        self.assertIn("attachment", disposition)
        self.assertIn("%E6%88%90%E7%BB%A9-", disposition)   # 成绩-（RFC 5987 转义）
        self.assertIn(".csv", disposition)
        self.assertTrue(data)

    def test_no_attempts_yet_is_an_empty_but_valid_export(self):
        """没有成绩的卷：导出应当是「只有表头」的成功响应，而不是 404/500。"""
        paper = _paper(self.server, title="M5 成绩验收卷-G（零成绩）", count=3)
        response = self.server.export_response(format="csv", paper_id=paper["id"])
        self.assertEqual(200, response.status_code, response.text)
        header, records = m5.csv_records(response.content)
        problems, rows = m5.csv_problems(data=response.content, expected_rows=0)
        m5.write_evidence("task42-export-empty.json", {
            "status": response.status_code, "bytes": len(response.content),
            "bom": response.content[:3].hex(), "header": header,
            "records": len(records), "problems": problems,
        })
        self.assertEqual(m5.BOM, response.content[:3])
        self.assertEqual(list(m5.COLUMN_LABELS), header)
        self.assertEqual([], records)
        self.assertEqual([], rows)
        self.assertEqual([], problems)


if __name__ == "__main__":
    unittest.main()
