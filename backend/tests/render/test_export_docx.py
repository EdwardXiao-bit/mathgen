"""Word（.docx）导出测试 —— T-M3-DOCX（M3）。

覆盖任务卡验收标准：

* docx 能生成、是合法 zip（`PK` 魔数）且含 `word/document.xml`
* 能被 `python-docx` 读回；题目是**真正的文字**（可被 Word 搜索到），不是图片
* `mode=QUESTION` 不含 answer / steps / difficulty / tags（安全边界）
* **四个 `answer_layout` 开关逐个生效**（关掉的那项不得出现）
* `cols_per_row` / `answer_line` / `header_bar` / `page_margin` 生效
* 导出确定性（同一 ctx 两次导出正文等价）
* HTML 与 DOCX **结构一致**（契约第 6 条）：同一份 plan、同样数量的答案条目
* LibreOffice 打开检查：docx -> PDF 成功且页数合理（工具缺失时报环境限制，不静默跳过）
"""
from __future__ import annotations

import os
import re
import subprocess
import tempfile
import unittest
import zipfile

from docx import Document
from docx.shared import Cm

from app.contracts.api import AnswerLayoutIn
from app.contracts.layout import LayoutParams
from app.contracts.sheet import SheetContext, SheetItem, SheetMode
from app.export import DocxExportError, docx_body_text, export_docx
from app.layout import build_answer_plan
from app.render import render_sheet

from tests.render.fixtures import build_ctx

BACKEND = os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
ARTIFACTS = os.path.join(BACKEND, ".tmp-render-artifacts")

#: 标记文本：用于「关掉那项就不得出现」的判定
STEP_MARK = "STEPMARK"
TAG_MARK = "TAGMARK"
DIFFICULTY_LABEL = "难度"

#: 打包进来的 LibreOffice Kit（AGENTS.md 记录的调用方式；不搜索系统 soffice）
LO_NODE = os.environ.get(
    "MATHGEN_LIBREOFFICE_NODE",
    r"C:\Users\Ed\AppData\Local\Programs\DeepSeek Harness\resources\runtime\primary-runtime\dependencies\node\bin\node.exe",
)
LO_CLI = os.environ.get(
    "MATHGEN_LIBREOFFICE_CLI",
    r"C:\Users\Ed\AppData\Local\Programs\DeepSeek Harness\resources\app.asar.unpacked\dsh\node_modules\@deepseek-ai\libreoffice-kit\lib\cli.js",
)


def _artifact(name: str) -> str:
    os.makedirs(ARTIFACTS, exist_ok=True)
    return os.path.join(ARTIFACTS, name)


def answer_ctx(
    count: int = 6,
    layout: LayoutParams | None = None,
    **answer_layout,
) -> SheetContext:
    """答案版卷面：每个可开关项都带可搜索的标记文本。"""
    items = tuple(
        SheetItem(
            seq=index + 1,
            display=f"{120 + index} × {12 + index} =",
            answer=f"ANS{index + 1}",
            difficulty=(index % 5) + 1,
            tags=(f"{TAG_MARK}{index + 1}",),
            steps=(f"{STEP_MARK}-A{index + 1}", f"{STEP_MARK}-B{index + 1}"),
        )
        for index in range(count)
    )
    return SheetContext(
        title="DOCX 测试卷",
        mode=SheetMode.ANSWER,
        items=items,
        layout=layout or LayoutParams(),
        code="2610-0001",
        header_note="限时 10 分钟",
        answer_layout=AnswerLayoutIn(**answer_layout),
    )


def question_ctx(count: int = 6, layout: LayoutParams | None = None) -> SheetContext:
    items = tuple(
        SheetItem(
            seq=index + 1,
            display=f"SECRETDISPLAY{index + 1}",
            answer=f"SECRETANSWER{index + 1}",
            difficulty=3,
            tags=(f"SECRETTAG{index + 1}",),
            steps=(f"SECRETSTEP{index + 1}",),
        )
        for index in range(count)
    )
    return SheetContext(
        title="题目版",
        mode=SheetMode.QUESTION,
        items=items,
        layout=layout or LayoutParams(),
    )


def export_text(ctx: SheetContext, name: str) -> str:
    export_docx(ctx, _artifact(name))
    return docx_body_text(_artifact(name))


class TestDocxFileShape(unittest.TestCase):
    """产物是合法 OOXML，且能被 python-docx 读回。"""

    def test_file_is_created_and_is_a_zip_with_pk_magic(self):
        ctx = answer_ctx()
        path = _artifact("shape.docx")
        export_docx(ctx, path)
        self.assertTrue(os.path.exists(path))
        with open(path, "rb") as handle:
            self.assertEqual(handle.read(2), b"PK", "docx 必须是 zip（PK 魔数）")
        self.assertTrue(zipfile.is_zipfile(path))

    def test_archive_contains_word_document_xml(self):
        ctx = answer_ctx()
        path = _artifact("parts.docx")
        export_docx(ctx, path)
        with zipfile.ZipFile(path) as archive:
            names = set(archive.namelist())
        self.assertIn("word/document.xml", names)
        self.assertIn("[Content_Types].xml", names)

    def test_readable_by_python_docx(self):
        ctx = answer_ctx()
        path = _artifact("readback.docx")
        export_docx(ctx, path)
        document = Document(path)
        self.assertGreater(len(document.paragraphs), 0)
        self.assertGreater(len(document.sections), 0)

    def test_question_text_is_real_text_not_an_image(self):
        """题目必须是**文字**（Word 里可搜索），不能是图片。"""
        ctx = answer_ctx()
        path = _artifact("text.docx")
        export_docx(ctx, path)
        with zipfile.ZipFile(path) as archive:
            media = [n for n in archive.namelist() if n.startswith("word/media/")]
        self.assertEqual(media, [], "卷面不应当包含图片（Word 里必须可编辑文字）")
        text = docx_body_text(path)
        for item in ctx.items:
            with self.subTest(seq=item.seq):
                self.assertIn(item.display, text)

    def test_rejects_non_context(self):
        with self.assertRaises(DocxExportError):
            export_docx("not a ctx", _artifact("bad.docx"))  # type: ignore[arg-type]


class TestPageSetup(unittest.TestCase):
    """A4 纸张 + `layout.page_margin` 页边距。"""

    def test_a4_portrait_with_page_margin(self):
        for margin in (1.0, 2.0, 3.0):
            ctx = answer_ctx(layout=LayoutParams(page_margin=margin))
            path = _artifact(f"page-{margin}.docx")
            export_docx(ctx, path)
            section = Document(path).sections[0]
            with self.subTest(margin=margin):
                self.assertAlmostEqual(section.page_width.cm, 21.0, delta=0.05)
                self.assertAlmostEqual(section.page_height.cm, 29.7, delta=0.05)
                self.assertAlmostEqual(section.left_margin.cm, margin, delta=0.05)
                self.assertAlmostEqual(section.right_margin.cm, margin, delta=0.05)
                self.assertAlmostEqual(section.top_margin.cm, margin, delta=0.05)
                self.assertAlmostEqual(section.bottom_margin.cm, margin, delta=0.05)


class TestQuestionModeSafety(unittest.TestCase):
    """契约第 3 条：`mode=QUESTION` 不得泄露 answer / steps / difficulty / tags。"""

    def test_no_answer_or_step_or_tag_text(self):
        text = export_text(question_ctx(), "safety.docx")
        for token in ("SECRETANSWER", "SECRETSTEP", "SECRETTAG", "难度"):
            with self.subTest(token=token):
                self.assertNotIn(token, text)

    def test_no_answer_section_heading(self):
        text = export_text(question_ctx(), "safety-heading.docx")
        self.assertNotIn("答案与解析", text)

    def test_question_display_is_present(self):
        text = export_text(question_ctx(), "safety-display.docx")
        for index in range(6):
            with self.subTest(index=index):
                self.assertIn(f"SECRETDISPLAY{index + 1}", text)

    def test_plan_for_question_mode_is_disabled(self):
        plan = build_answer_plan(question_ctx())
        self.assertFalse(plan.enabled)
        self.assertEqual(plan.entries, ())


class TestAnswerLayoutSwitches(unittest.TestCase):
    """四个 `answer_layout` 开关**逐个**验证：关掉哪项，那项就不得出现。"""

    def setUp(self):
        self.baseline = export_text(answer_ctx(), "switches-all-on.docx")

    def test_baseline_contains_everything(self):
        for token in ("ANS1", DIFFICULTY_LABEL, TAG_MARK, STEP_MARK):
            with self.subTest(token=token):
                self.assertIn(token, self.baseline)

    def test_include_steps_false_removes_steps(self):
        text = export_text(answer_ctx(include_steps=False), "switches-steps-off.docx")
        self.assertNotIn(STEP_MARK, text, "include_steps=False 时不得出现逐步过程")
        self.assertIn(DIFFICULTY_LABEL, text, "其它项不应被连带关掉")
        self.assertIn(TAG_MARK, text, "其它项不应被连带关掉")
        self.assertIn("ANS1", text)

    def test_include_difficulty_false_removes_difficulty(self):
        text = export_text(answer_ctx(include_difficulty=False), "switches-diff-off.docx")
        self.assertNotIn(DIFFICULTY_LABEL, text, "include_difficulty=False 时不得出现难度")
        self.assertIn(STEP_MARK, text)
        self.assertIn(TAG_MARK, text)
        self.assertIn("ANS1", text)

    def test_include_tags_false_removes_tags(self):
        text = export_text(answer_ctx(include_tags=False), "switches-tags-off.docx")
        self.assertNotIn(TAG_MARK, text, "include_tags=False 时不得出现考点标签")
        self.assertIn(STEP_MARK, text)
        self.assertIn(DIFFICULTY_LABEL, text)
        self.assertIn("ANS1", text)

    def test_separate_page_false_removes_page_break_before_answers(self):
        """`separate_page=False` 时答案区前**不得有分页符**。

        判定方式：数 `word/document.xml` 里的 `<w:br w:type="page"/>`。
        """
        on = _artifact("sep-on.docx")
        off = _artifact("sep-off.docx")
        export_docx(answer_ctx(separate_page=True), on)
        export_docx(answer_ctx(separate_page=False), off)

        def page_breaks(path: str) -> int:
            with zipfile.ZipFile(path) as archive:
                xml = archive.read("word/document.xml").decode("utf-8")
            return len(re.findall(r'w:type="page"', xml))

        self.assertGreaterEqual(page_breaks(on), 1, "separate_page=True 时应有分页符")
        self.assertEqual(page_breaks(off), 0, "separate_page=False 时不应有分页符")

    def test_all_switches_off_keeps_only_answers(self):
        text = export_text(
            answer_ctx(include_steps=False, include_difficulty=False, include_tags=False),
            "switches-all-off.docx",
        )
        for token in (STEP_MARK, TAG_MARK, DIFFICULTY_LABEL):
            with self.subTest(token=token):
                self.assertNotIn(token, text)
        self.assertIn("ANS1", text, "答案本身始终要有")

    def test_each_switch_is_independently_effective(self):
        """四个开关两两独立：任一单项关闭只影响它自己。"""
        cases = {
            "include_steps": STEP_MARK,
            "include_difficulty": DIFFICULTY_LABEL,
            "include_tags": TAG_MARK,
        }
        for field, token in cases.items():
            with self.subTest(field=field):
                text = export_text(answer_ctx(**{field: False}), f"indep-{field}.docx")
                self.assertNotIn(token, text)
                others = {k: v for k, v in cases.items() if k != field}
                for other_token in others.values():
                    self.assertIn(other_token, text, f"关掉 {field} 不应影响 {other_token}")

    def test_answer_layout_also_drives_html(self):
        """契约第 4 条：HTML 也要尊重同一组开关（与 DOCX 同一条通路）。"""
        html = render_sheet(answer_ctx(include_steps=False, include_tags=False))
        body = html.split("</style>", 1)[-1]
        self.assertNotIn(STEP_MARK, body)
        self.assertNotIn(TAG_MARK, body)
        self.assertIn("ANS1", body)


class TestLayoutParity(unittest.TestCase):
    """`cols_per_row` / `answer_line` / `header_bar` 在 docx 里生效。"""

    def test_cols_per_row_sets_table_columns(self):
        for cols in (2, 3, 4, 6):
            ctx = answer_ctx(count=6, layout=LayoutParams(cols_per_row=cols))
            path = _artifact(f"cols-{cols}.docx")
            export_docx(ctx, path)
            document = Document(path)
            with self.subTest(cols=cols):
                self.assertTrue(document.tables, "题目区应当用表格排布")
                question_table = document.tables[-1]
                self.assertEqual(len(question_table.columns), cols)

    def test_answer_line_controls_blank_underline(self):
        on = _artifact("aline-on.docx")
        off = _artifact("aline-off.docx")
        export_docx(answer_ctx(layout=LayoutParams(answer_line=True)), on)
        export_docx(answer_ctx(layout=LayoutParams(answer_line=False)), off)

        def borders(path: str) -> int:
            with zipfile.ZipFile(path) as archive:
                xml = archive.read("word/document.xml").decode("utf-8")
            return xml.count("<w:bottom")

        self.assertGreater(borders(on), borders(off), "answer_line=True 时答题横线更多")

    def test_header_bar_controls_header_paragraph(self):
        on_path = _artifact("hbar-on.docx")
        off_path = _artifact("hbar-off.docx")
        export_docx(answer_ctx(layout=LayoutParams(header_bar=True)), on_path)
        export_docx(answer_ctx(layout=LayoutParams(header_bar=False)), off_path)
        self.assertIn("班级", docx_body_text(on_path))
        self.assertNotIn("班级", docx_body_text(off_path), "header_bar=False 时不得出现卷头栏")

    def test_title_code_and_note_are_exported(self):
        text = export_text(answer_ctx(), "meta.docx")
        self.assertIn("DOCX 测试卷", text)
        self.assertIn("2610-0001", text)
        self.assertIn("限时 10 分钟", text)

    def test_question_table_row_count_matches_chunking(self):
        from app.render.sheet import chunk_rows

        for count, cols in ((6, 2), (6, 3), (7, 4), (12, 6)):
            ctx = answer_ctx(count=count, layout=LayoutParams(cols_per_row=cols))
            path = _artifact(f"rows-{count}-{cols}.docx")
            export_docx(ctx, path)
            expected = len(chunk_rows(tuple(ctx.items), cols))
            with self.subTest(count=count, cols=cols):
                self.assertEqual(len(Document(path).tables[-1].rows), expected)

    def test_empty_sheet_exports_with_placeholder(self):
        ctx = SheetContext(title="空卷", mode=SheetMode.QUESTION, items=(),
                           layout=LayoutParams())
        text = export_text(ctx, "empty.docx")
        self.assertIn("本卷暂无题目", text)


class TestDeterminism(unittest.TestCase):
    """同一 `SheetContext` 导出两次，内容等价（时间戳等易变字段不得写进文档）。"""

    def test_body_text_is_identical_across_exports(self):
        ctx = answer_ctx()
        first = _artifact("det-1.docx")
        second = _artifact("det-2.docx")
        export_docx(ctx, first)
        export_docx(ctx, second)
        self.assertEqual(docx_body_text(first), docx_body_text(second))

    def test_document_xml_has_no_timestamp_or_author(self):
        path = _artifact("det-nometa.docx")
        export_docx(answer_ctx(), path)
        with zipfile.ZipFile(path) as archive:
            xml = archive.read("word/document.xml").decode("utf-8")
            core = archive.read("docProps/core.xml").decode("utf-8")
        self.assertNotIn("dcterms:created", xml)
        # 作者字段被显式清空（不把「谁导的」写进交付物）
        self.assertIn("<dc:creator></dc:creator>", core.replace(" ", ""))

    def test_equal_contexts_produce_identical_text(self):
        self.assertEqual(
            export_text(answer_ctx(), "det-a.docx"),
            export_text(answer_ctx(), "det-b.docx"),
        )


class TestHtmlDocxParity(unittest.TestCase):
    """契约第 6 条：DOCX 与 HTML 结构一致（同一份 plan）。"""

    def test_same_answer_entry_count(self):
        for count in (1, 4, 9):
            ctx = answer_ctx(count=count)
            html = render_sheet(ctx)
            body = html.split("</style>", 1)[-1]
            path = _artifact(f"parity-{count}.docx")
            export_docx(ctx, path)
            text = docx_body_text(path)
            with self.subTest(count=count):
                self.assertEqual(body.count('<div class="a-item"'), count)
                self.assertEqual(sum(text.count(f"ANS{i + 1}") for i in range(count)), count)

    def test_answer_entries_are_in_same_order(self):
        ctx = answer_ctx(count=5)
        path = _artifact("parity-order.docx")
        export_docx(ctx, path)
        text = docx_body_text(path)
        positions = [text.index(f"ANS{i + 1}") for i in range(5)]
        self.assertEqual(positions, sorted(positions), "答案顺序应与卷面题号一致")

    def test_question_display_order_is_preserved(self):
        ctx = answer_ctx(count=5)
        path = _artifact("parity-qorder.docx")
        export_docx(ctx, path)
        text = docx_body_text(path)
        positions = [text.index(item.display) for item in ctx.items]
        self.assertEqual(positions, sorted(positions), "题号顺序不得被重排")


class TestLibreOfficeOpen(unittest.TestCase):
    """**必须真的用 LibreOffice 打开检查**：docx -> PDF。"""

    @classmethod
    def setUpClass(cls):
        cls.skip_reason = None
        if not os.path.exists(LO_NODE) or not os.path.exists(LO_CLI):
            cls.skip_reason = (
                "打包的 LibreOffice Kit 不可用（这是**环境限制**，不是实现问题）："
                f"node={LO_NODE!r} cli={LO_CLI!r}；"
                "结构检查仍由本文件其余用例覆盖。"
            )

    def _convert(self, docx_path: str) -> str:
        if self.skip_reason:
            self.skipTest(self.skip_reason)
        pdf_path = docx_path[:-5] + ".pdf"
        if os.path.exists(pdf_path):
            os.remove(pdf_path)          # 输出必须是新文件

        # 本机实测坑（与 Playwright 同源但不同进程）：LibreOffice Kit 会在**系统 TEMP**
        # 下 `mkdtemp('libreoffice-kit-XXXXXX')`，沙箱里得到
        # `EPERM: operation not permitted`。对策：给子进程传一个可写的 TMP/TEMP。
        # 用 `tempfile.mkdtemp` 在仓库内建目录，避免把本机路径硬编进测试。
        temp_dir = tempfile.mkdtemp(prefix="lo-kit-", dir=_artifact("."))
        env = dict(os.environ)
        env["TMP"] = env["TEMP"] = env["TMPDIR"] = temp_dir

        result = subprocess.run(
            [LO_NODE, LO_CLI, "convert", "--input", docx_path, "--output", pdf_path],
            capture_output=True,
            text=True,
            encoding="utf-8",
            errors="replace",          # 转换器在中文 Windows 上会输出非 UTF-8 字节
            timeout=300,
            env=env,
        )
        self.assertEqual(result.returncode, 0, f"LibreOffice 转换失败：{result.stderr[:500]}")
        self.assertTrue(os.path.exists(pdf_path), "转换命令成功但没有产物")
        return pdf_path

    def test_question_sheet_converts_to_pdf(self):
        from app.export import count_pdf_pages

        path = _artifact("lo-question.docx")
        export_docx(question_ctx(count=24), path)
        pdf_path = self._convert(path)
        with open(pdf_path, "rb") as handle:
            pages = count_pdf_pages(handle.read())
        self.assertGreaterEqual(pages, 1, "LibreOffice 转出的 PDF 应当至少 1 页")
        self.assertLess(pages, 20, f"24 道题不应超过 20 页，实际 {pages} 页")
        self.assertGreater(os.path.getsize(pdf_path), 1000)

    def test_answer_sheet_converts_to_pdf_with_more_pages(self):
        from app.export import count_pdf_pages

        question = _artifact("lo-answer-q.docx")
        answer = _artifact("lo-answer-a.docx")
        export_docx(question_ctx(count=24), question)
        export_docx(answer_ctx(count=24), answer)

        def pages(path: str) -> int:
            pdf = self._convert(path)
            with open(pdf, "rb") as handle:
                return count_pdf_pages(handle.read())

        question_pages = pages(question)
        answer_pages = pages(answer)
        self.assertGreaterEqual(answer_pages, question_pages)

    def test_separate_page_false_fits_on_fewer_pages(self):
        """`separate_page=False` 时答案不另起一页 —— 页数不应增加。"""
        from app.export import count_pdf_pages

        on = _artifact("lo-sep-on.docx")
        off = _artifact("lo-sep-off.docx")
        export_docx(answer_ctx(count=6, separate_page=True), on)
        export_docx(answer_ctx(count=6, separate_page=False), off)

        def pages(path: str) -> int:
            pdf = self._convert(path)
            with open(pdf, "rb") as handle:
                return count_pdf_pages(handle.read())

        self.assertLessEqual(pages(off), pages(on))


if __name__ == "__main__":
    unittest.main()
