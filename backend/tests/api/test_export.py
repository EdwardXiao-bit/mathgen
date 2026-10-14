"""`GET /api/papers/{id}/export` 与「未匹配路由统一 404」测试（契约 §5.1 / §5.4，D57）。

分两层：

* **默认层（不依赖浏览器）**：把 `pdf_export.export_pdf_bytes` 换成探针，
  断言路由装配、状态码、统一错误结构与 **C3 同源**（导出用的 HTML 与 `/sheet`
  逐字节相同）。这一层在任何机器上都必须过。
* **实物层（需要 Chromium）**：真的走完 `render_sheet → export_pdf`，
  断言 `%PDF` 魔数、`application/pdf`、答案版页数更多、连续导出可复现。
  浏览器不可用时 `skip`（环境问题，不是实现问题），实物证据另见交接报告。
"""
from __future__ import annotations

import os
import sys
import unittest
from pathlib import Path

# 本机浏览器装在仓库内 `.playwright`（见 AGENTS.md「坑 2」）；其他机器走默认路径
_BACKEND = Path(__file__).resolve().parents[2]
_REPO = _BACKEND.parent
if (_REPO / ".playwright").is_dir():
    os.environ.setdefault("PLAYWRIGHT_BROWSERS_PATH", str(_REPO / ".playwright"))

from app.contracts.api import ROUTES  # noqa: E402
from app.services import pdf_export  # noqa: E402
from tests.api._support import (  # noqa: E402
    ApiTestCase,
    docx_unavailable,
    docx_with,
    render_unavailable,
    url_of,
)

#: Word 导出通道的 MIME（契约 §5.1）
MEDIA_TYPE_DOCX = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"


def _docx_module_available() -> bool:
    try:
        import app.export.docx  # noqa: F401
    except ImportError:
        return False
    return True


class _FakePdf:
    """伪导出器：记录收到的 HTML，返回一个最小的合法 PDF 头。"""

    def __init__(self, payload: bytes = b"%PDF-1.4\n% fake\n") -> None:
        self.payload = payload
        self.seen: list[str] = []

    def __call__(self, html: str) -> bytes:
        self.seen.append(html)
        return self.payload


class TestExportRouteHermetic(ApiTestCase):
    """不启动浏览器的路由装配与 C3 同源断言。"""

    def setUp(self) -> None:
        super().setUp()
        self.paper = self.create_paper({"count": 5})
        self.pid = self.paper["id"]
        self.fake = _FakePdf()
        self._real = pdf_export.export_pdf_bytes
        pdf_export.export_pdf_bytes = self.fake

    def tearDown(self) -> None:
        pdf_export.export_pdf_bytes = self._real
        super().tearDown()

    # ---- 正常路径 -------------------------------------------------------

    def test_pdf_export_returns_pdf_stream(self):
        response = self.client.get(
            url_of("export", paper_id=self.pid), params={"format": "pdf", "mode": "question"}
        )
        self.assertEqual(response.status_code, 200, response.text)
        self.assertEqual(response.headers["content-type"], "application/pdf")
        self.assertTrue(response.content.startswith(b"%PDF"))
        self.assertEqual(
            response.headers["content-disposition"],
            f'attachment; filename="paper-{self.pid}-question.pdf"',
        )
        self.assertEqual(len(self.fake.seen), 1)

    def test_export_mode_defaults_to_question(self):
        response = self.client.get(url_of("export", paper_id=self.pid))
        self.assertEqual(response.status_code, 200, response.text)
        self.assertIn("data-mode=\"question\"", self.fake.seen[0])

    def test_export_format_defaults_to_pdf(self):
        response = self.client.get(url_of("export", paper_id=self.pid), params={"mode": "answer"})
        self.assertEqual(response.status_code, 200, response.text)
        self.assertTrue(response.content.startswith(b"%PDF"))

    def test_export_format_is_case_insensitive(self):
        response = self.client.get(
            url_of("export", paper_id=self.pid), params={"format": "PDF"}
        )
        self.assertEqual(response.status_code, 200, response.text)
        self.assertTrue(response.content.startswith(b"%PDF"))

    def test_export_html_is_byte_identical_to_sheet(self):
        """C3 硬证据：`/export` 打印的 HTML 与 `/sheet` 返回的**逐字节相同**。"""
        for mode in ("question", "answer"):
            with self.subTest(mode=mode):
                self.fake.seen.clear()
                exported = self.client.get(
                    url_of("export", paper_id=self.pid), params={"format": "pdf", "mode": mode}
                )
                sheet = self.client.get(url_of("sheet", paper_id=self.pid), params={"mode": mode})
                self.assertEqual(exported.status_code, 200, exported.text)
                self.assertEqual(sheet.status_code, 200, sheet.text)
                self.assertEqual(
                    self.fake.seen,
                    [sheet.text],
                    "导出用的 HTML 必须与 /sheet 返回的完全一致（C3：同一份 HTML）",
                )

    def test_export_and_sheet_share_header_note(self):
        note = "限时 10 分钟"
        self.client.get(
            url_of("export", paper_id=self.pid),
            params={"format": "pdf", "mode": "answer", "header_note": note},
        )
        sheet = self.client.get(
            url_of("sheet", paper_id=self.pid), params={"mode": "answer", "header_note": note}
        )
        self.assertEqual(self.fake.seen, [sheet.text])
        self.assertIn(note, sheet.text)

    def test_export_works_for_draft_and_published(self):
        """导出不要求已发布：draft 也要能导出打印（与 /sheet 一致）。"""
        draft = self.pid
        self.assertEqual(
            self.client.get(url_of("export", paper_id=draft)).status_code, 200
        )
        self.publish(draft)
        self.assertEqual(
            self.client.get(url_of("export", paper_id=draft)).status_code, 200
        )
        self.assertEqual(len(self.fake.seen), 2)

    def test_export_sheet_and_export_after_edit_use_snapshot(self):
        """改题后再导出：导出的 HTML 必须反映**最新快照**（不重算引擎）。"""
        self.client.delete(url_of("delete_item", paper_id=self.pid, seq=1))
        self.fake.seen.clear()
        self.client.get(url_of("export", paper_id=self.pid))
        sheet = self.client.get(url_of("sheet", paper_id=self.pid))
        self.assertEqual(self.fake.seen, [sheet.text])
        self.assertEqual(sheet.text.count('class="q-item"'), len(self.paper["items"]) - 1)

    # ---- Word（M3 真实现） ---------------------------------------------

    def test_docx_export_returns_word_stream(self):
        """`format=docx` 走真实现：正确的 MIME 与 zip 魔数（render-export 的模块）。"""
        if not _docx_module_available():
            self.skipTest("app.export.docx 尚未实现（render-export 并行开发中）")
        captured: list = []

        def fake_export_docx(ctx, out_path):
            captured.append(ctx)
            with open(out_path, "wb") as handle:
                handle.write(b"PK\x03\x04fake-docx")

        with docx_with(fake_export_docx):
            response = self.client.get(
                url_of("export", paper_id=self.pid),
                params={"format": "docx", "mode": "answer"},
            )
        self.assertEqual(response.status_code, 200, response.text)
        self.assertEqual(response.headers["content-type"], MEDIA_TYPE_DOCX)
        self.assertTrue(response.content.startswith(b"PK\x03\x04"))
        self.assertEqual(
            response.headers["content-disposition"],
            f'attachment; filename="paper-{self.pid}-answer.docx"',
        )
        # Word 拿到的是与 /sheet 同一份 SheetContext（含 answer_layout）
        self.assertEqual(len(captured), 1)
        self.assertEqual(captured[0].mode.value, "answer")
        self.assertEqual([item.seq for item in captured[0].items], [1, 2, 3, 4, 5])

    def test_docx_ctx_carries_paper_answer_layout(self):
        """组装 Word 用的 SheetContext 必须带上落库的 `answer_layout`（M3 关键点）。"""
        self.client.patch(
            url_of("update_answer_layout", paper_id=self.pid),
            json={"separate_page": False, "include_steps": False,
                  "include_difficulty": True, "include_tags": False},
        )
        captured: list = []

        def fake_export_docx(ctx, out_path):
            captured.append(ctx)
            with open(out_path, "wb") as handle:
                handle.write(b"PK\x03\x04x")

        with docx_with(fake_export_docx):
            response = self.client.get(
                url_of("export", paper_id=self.pid), params={"format": "docx"}
            )
        self.assertEqual(response.status_code, 200, response.text)
        layout = captured[0].answer_layout
        self.assertFalse(layout.separate_page)
        self.assertFalse(layout.include_steps)
        self.assertTrue(layout.include_difficulty)
        self.assertFalse(layout.include_tags)

    def test_docx_503_when_module_missing(self):
        """render-export 还没就绪时必须是 503 `RENDER_NOT_READY`，不是 500。"""
        with docx_unavailable():
            response = self.client.get(
                url_of("export", paper_id=self.pid), params={"format": "docx"}
            )
        self.assertEqual(response.status_code, 503, response.text)
        body = response.json()
        self.assertEqual(body["error"]["code"], "RENDER_NOT_READY")
        self.assertIn("docx", body["error"]["message"])
        self.assertNotIn("detail", body)

    def test_docx_503_when_generator_fails(self):
        def broken(ctx, out_path):
            raise RuntimeError("模拟 Word 生成失败")

        with docx_with(broken):
            response = self.client.get(
                url_of("export", paper_id=self.pid), params={"format": "docx"}
            )
        self.assertEqual(response.status_code, 503, response.text)
        self.assertEqual(response.json()["error"]["code"], "EXPORT_NOT_READY")

    def test_docx_unknown_paper_is_404(self):
        response = self.client.get(url_of("export", paper_id=999999), params={"format": "docx"})
        self.assertEqual(response.status_code, 404, response.text)
        self.assertEqual(response.json()["error"]["code"], "PAPER_NOT_FOUND")

    # ---- 错误路径 -------------------------------------------------------

    def test_invalid_format_is_422(self):
        for value in ("doc", "xlsx", "", "PDFX"):
            with self.subTest(value=value):
                response = self.client.get(
                    url_of("export", paper_id=self.pid), params={"format": value}
                )
                self.assertEqual(response.status_code, 422, response.text)
                self.assertEqual(response.json()["error"]["code"], "VALIDATION_ERROR")

    def test_invalid_mode_is_422(self):
        response = self.client.get(
            url_of("export", paper_id=self.pid), params={"format": "pdf", "mode": "both"}
        )
        self.assertEqual(response.status_code, 422, response.text)
        self.assertEqual(response.json()["error"]["code"], "VALIDATION_ERROR")

    def test_unknown_paper_is_404(self):
        response = self.client.get(url_of("export", paper_id=999999))
        self.assertEqual(response.status_code, 404, response.text)
        self.assertEqual(response.json()["error"]["code"], "PAPER_NOT_FOUND")

    def test_unknown_paper_wins_over_bad_format(self):
        """错误优先级固定：卷不存在 → 404，即使 format 也非法。"""
        response = self.client.get(
            url_of("export", paper_id=999999), params={"format": "bogus"}
        )
        self.assertEqual(response.status_code, 404, response.text)
        self.assertEqual(response.json()["error"]["code"], "PAPER_NOT_FOUND")

    def test_export_503_when_render_not_ready(self):
        with render_unavailable():
            response = self.client.get(url_of("export", paper_id=self.pid))
        self.assertEqual(response.status_code, 503, response.text)
        self.assertEqual(response.json()["error"]["code"], "RENDER_NOT_READY")
        self.assertEqual(self.fake.seen, [])

    def test_export_503_when_browser_channel_unavailable(self):
        def broken(_html: str) -> bytes:
            raise pdf_export.ExportUnavailable("模拟：本机没有可用浏览器")

        pdf_export.export_pdf_bytes = broken
        response = self.client.get(url_of("export", paper_id=self.pid))
        self.assertEqual(response.status_code, 503, response.text)
        body = response.json()
        self.assertEqual(body["error"]["code"], "EXPORT_NOT_READY")
        self.assertIn("模拟", body["error"]["message"])
        self.assertNotIn("detail", body)


class TestUniformNotFound(ApiTestCase):
    """D57：未匹配路由 / 方法不允许也必须走统一错误结构，不泄漏 FastAPI 的 `detail`。"""

    def test_unknown_path_is_uniform_404(self):
        for path in ("/api/nonexistent", "/api/papers/1/nope", "/nope"):
            with self.subTest(path=path):
                response = self.client.get(path)
                self.assertEqual(response.status_code, 404, response.text)
                body = response.json()
                self.assertEqual(set(body), {"error"}, body)
                self.assertEqual(set(body["error"]), {"code", "message"})
                self.assertEqual(body["error"]["code"], "NOT_FOUND")
                self.assertNotIn("detail", body)
                self.assertTrue(body["error"]["message"].strip())

    def test_unknown_path_message_names_the_route(self):
        body = self.client.get("/api/nonexistent").json()
        self.assertIn("/api/nonexistent", body["error"]["message"])

    def test_wrong_method_is_uniform_405(self):
        created = self.create_paper({"count": 2})
        response = self.client.get(url_of("publish_paper", paper_id=created["id"]))
        self.assertEqual(response.status_code, 405, response.text)
        body = response.json()
        self.assertEqual(body["error"]["code"], "METHOD_NOT_ALLOWED")
        self.assertNotIn("detail", body)

    def test_unknown_path_404_does_not_leak_detail_for_post(self):
        response = self.client.post("/api/nope", json={})
        self.assertEqual(response.status_code, 404, response.text)
        self.assertNotIn("detail", response.json())


@unittest.skipUnless(
    _browser_available := __import__("app.export", fromlist=["is_available"]).is_available(),
    "本机没有可用 Chromium，实物 PDF 用例跳过（环境问题，见 AGENTS.md「坑 2」）",
)
class TestExportPdfReal(ApiTestCase):
    """真实 PDF 实物：必须走完整链路（render_sheet → 子进程 export_pdf）。"""

    @classmethod
    def setUpClass(cls) -> None:
        super().setUpClass()
        from app.export import count_pdf_pages

        cls.count_pdf_pages = staticmethod(count_pdf_pages)

    def setUp(self) -> None:
        super().setUp()
        self.paper = self.create_paper({"count": 8})
        self.pid = self.paper["id"]

    def _export(self, mode: str):
        return self.client.get(
            url_of("export", paper_id=self.pid), params={"format": "pdf", "mode": mode}
        )

    def test_real_pdf_magic_and_page_counts(self):
        """题目版与答案版都出真 PDF，且答案版页数更多（D20 答案单独成页）。"""
        question = self._export("question")
        self.assertEqual(question.status_code, 200, question.text[:400])
        self.assertTrue(question.content.startswith(b"%PDF"))
        self.assertEqual(question.headers["content-type"], "application/pdf")

        answer = self._export("answer")
        self.assertEqual(answer.status_code, 200, answer.text[:400])
        self.assertTrue(answer.content.startswith(b"%PDF"))

        question_pages = self.count_pdf_pages(question.content)
        answer_pages = self.count_pdf_pages(answer.content)
        self.assertGreaterEqual(question_pages, 1)
        self.assertGreater(answer_pages, question_pages,
                           f"答案版应多出答案页（question={question_pages}, answer={answer_pages}）")

    def test_real_pdf_is_reproducible_across_requests(self):
        """连续两次导出（FastAPI 线程池里可能是不同线程）都成功且字节相同。"""
        first = self._export("question")
        second = self._export("question")
        self.assertEqual(first.status_code, 200, first.text[:400])
        self.assertEqual(second.status_code, 200, second.text[:400])
        self.assertEqual(first.content, second.content,
                         "同一份 HTML 两次导出应逐字节相同（export_pdf 已归一化时间）")


@unittest.skipUnless(
    _docx_module_available(),
    "app.export.docx 尚未实现（render-export 并行开发中）；503 路径已由其他用例覆盖",
)
class TestExportDocxReal(ApiTestCase):
    """真实 Word 实物：必须走完整链路（build_sheet_context → export_docx）。"""

    def _export(self, paper_id: int, mode: str):
        return self.client.get(
            url_of("export", paper_id=paper_id), params={"format": "docx", "mode": mode}
        )

    def test_real_docx_end_to_end(self):
        created = self.create_paper({"count": 4}, title="Word 实物卷")
        question = self._export(created["id"], "question")
        answer = self._export(created["id"], "answer")

        for response in (question, answer):
            self.assertEqual(response.status_code, 200, response.text[:300])
            self.assertEqual(response.headers["content-type"], MEDIA_TYPE_DOCX)
            self.assertTrue(response.content.startswith(b"PK\x03\x04"), "必须是 zip 容器")

        question_text = _docx_text(question.content)
        answer_text = _docx_text(answer.content)

        # 卷头与题干都在（题目在表格里，不能只看 paragraphs）
        self.assertIn(created["title"], question_text)
        self.assertIn(created["items"][0]["display"], question_text)
        self.assertIn("班级", question_text)

        # 题目版不得出现答案区；答案版必须有
        self.assertNotIn("答案与解析", question_text)
        self.assertIn("答案与解析", answer_text)
        self.assertIn(created["items"][0]["answer"], answer_text)

    def test_real_docx_is_reproducible(self):
        created = self.create_paper({"count": 3})
        first = self._export(created["id"], "question")
        second = self._export(created["id"], "question")
        self.assertEqual(first.status_code, 200, first.text[:300])
        self.assertEqual(second.status_code, 200, second.text[:300])
        self.assertEqual(_docx_text(first.content), _docx_text(second.content),
                         "同一份卷面两次导出的正文应一致")

    def test_docx_reflects_answer_layout_edit(self):
        """改答案版形式后导出的 Word 必须随之变化（同源的真实证据）。"""
        created = self.create_paper({"count": 3})
        with_tags = self._export(created["id"], "answer").content
        self.assertIn(created["items"][0]["tags"][0], _docx_text(with_tags))

        self.client.patch(
            url_of("update_answer_layout", paper_id=created["id"]),
            json={"separate_page": True, "include_steps": True,
                  "include_difficulty": False, "include_tags": False},
        )
        without_tags = self._export(created["id"], "answer").content
        text = _docx_text(without_tags)
        self.assertIn("答案与解析", text)
        self.assertNotIn("难度", text, "include_difficulty=False 时不应标难度")


def _docx_text(payload: bytes) -> str:
    """把 .docx 的段落与表格单元格拼成纯文本（用于内容断言）。"""
    import io

    import docx

    document = docx.Document(io.BytesIO(payload))
    chunks = [paragraph.text for paragraph in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            chunks.extend(cell.text for cell in row.cells)
    return "\n".join(chunk for chunk in chunks if chunk)


if __name__ == "__main__":  # pragma: no cover
    unittest.main()
