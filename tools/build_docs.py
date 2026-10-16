# -*- coding: utf-8 -*-
"""把 docs/*.md 生成为课程提交用的 Word 文档（D36 / D37）。

设计原则：
  - md 是唯一事实来源，本脚本只做"派生"，绝不反向修改 md。
  - 图的位置由 ANCHORS 显式声明；锚点找不到就报错退出，避免静默错位。
  - 表自动编号并加题注（题注取所属小节的标题），图同理由 ANCHORS 给定。
  - 宽表（fit_tables=True）按内容分配列宽并禁止行被分页切开；默认路径不启用，
    保证既有产物字节不变。
"""
import os
import re
import sys

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
DOCS = os.path.join(ROOT, "docs")
FIGS = os.path.join(DOCS, "word", "figures")
OUTDIR = os.path.join(DOCS, "word")

SONG = "宋体"
HEI = "黑体"
MONO = "Consolas"

COVER = {
    "course": "现代软件工程课程设计",
    "project": "小学四则运算练习题生成与批改系统",
    "members": "肖铭皓    胡宇书",
    "sid": "肖：待填　胡：待填",
    "teacher": "罗以宁",
}

# 完成日期**按文档**给定（原来是全局常量，M7 加第三份文档时改成按文档传参）：
# 01/02 在 2026-10-14 交付；03《测试文档》要等 M6 收口、有真实执行记录之后才写。
DATES = {
    "01-需求分析文档.docx": "2026 年 10 月 14 日",
    "02-技术设计文档.docx": "2026 年 10 月 14 日",
    "03-测试文档.docx": "2026 年 10 月 16 日",
}


# ============================================================ markdown 解析
def split_row(line):
    r"""按 `|` 切一行表格；`\|` 是 markdown 的转义竖线，必须留在单元格里。

    docs/02 的表里有 `` `...?format=csv\|xlsx` `` 这种写法（4 行）。按裸 `|` 切会把一个
    行内代码劈成两半，半个反引号落进前一个单元格、半个落进后一个单元格，单元格数还比
    表头多一个——Word 里既看到裸反引号，又看到错行的表格。
    """
    line = line.strip()
    if line.startswith("|"):
        line = line[1:]
    if line.endswith("|") and not line.endswith("\\|"):
        line = line[:-1]
    cells, cur, i = [], [], 0
    while i < len(line):
        ch = line[i]
        if ch == "\\" and i + 1 < len(line) and line[i + 1] == "|":
            cur.append("|")
            i += 2
            continue
        if ch == "|":
            cells.append("".join(cur))
            cur = []
            i += 1
            continue
        cur.append(ch)
        i += 1
    cells.append("".join(cur))
    return [c.strip() for c in cells]


def parse_md(path):
    lines = open(path, encoding="utf-8").read().split("\n")
    blocks, i = [], 0
    while i < len(lines):
        ln = lines[i]
        if ln.startswith("```"):
            buf, i = [], i + 1
            while i < len(lines) and not lines[i].startswith("```"):
                buf.append(lines[i])
                i += 1
            i += 1
            blocks.append(("code", "\n".join(buf)))
            continue
        m = re.match(r"^(#{1,6})\s+(.*)$", ln)
        if m:
            blocks.append(("h", len(m.group(1)), m.group(2).strip()))
            i += 1
            continue
        if ln.strip().startswith("|") and i + 1 < len(lines) and \
                re.match(r"^\|[\s:\-|]+\|$", lines[i + 1].strip()):
            header = split_row(ln)
            i += 2
            rows = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                rows.append(split_row(lines[i]))
                i += 1
            blocks.append(("table", header, rows))
            continue
        if ln.strip().startswith(">"):
            buf = []
            while i < len(lines) and lines[i].strip().startswith(">"):
                buf.append(lines[i].strip().lstrip(">").strip())
                i += 1
            blocks.append(("quote", " ".join(x for x in buf if x)))
            continue
        if re.match(r"^\s*[-*]\s+", ln):
            items = []
            while i < len(lines) and re.match(r"^\s*[-*]\s+", lines[i]):
                items.append(re.sub(r"^\s*[-*]\s+", "", lines[i]))
                i += 1
            blocks.append(("ul", items))
            continue
        if re.match(r"^\s*\d+\.\s+", ln):
            items = []
            while i < len(lines) and re.match(r"^\s*\d+\.\s+", lines[i]):
                items.append(re.sub(r"^\s*\d+\.\s+", "", lines[i]))
                i += 1
            blocks.append(("ol", items))
            continue
        if ln.strip() == "---":
            blocks.append(("hr",))
            i += 1
            continue
        if not ln.strip():
            i += 1
            continue
        buf = [ln]
        i += 1
        while i < len(lines) and lines[i].strip() and not re.match(
                r"^(#{1,6}\s|>|\s*[-*]\s|\s*\d+\.\s|\||```|---)", lines[i]):
            buf.append(lines[i])
            i += 1
        blocks.append(("p", " ".join(buf)))
    return blocks


# ============================================================ docx 基础件
def set_font(run, ascii_font, ea_font, size, bold=False, color=None, mono=False):
    run.font.size = Pt(size)
    run.font.bold = bold
    if color:
        run.font.color.rgb = color
    name = MONO if mono else ascii_font
    run.font.name = name
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    rfonts.set(qn("w:ascii"), name)
    rfonts.set(qn("w:hAnsi"), name)
    rfonts.set(qn("w:eastAsia"), ea_font)


def add_inline(par, text, size=12, ea=SONG, ascii_font="Times New Roman", bold=False):
    text = text.replace("<br>", " ").replace("<br/>", " ")
    for seg in re.split(r"(\*\*.+?\*\*|`[^`]+`)", text):
        if not seg:
            continue
        if seg.startswith("**") and seg.endswith("**") and len(seg) > 4:
            # 粗体段内部可能还嵌着 `代码`（如 `**为什么用 \`unittest\` 而不是 pytest**`）：
            # 再切一层，让代码段也走 MONO 分支，否则反引号会被原样写进 Word。
            for inner in re.split(r"(`[^`]+`)", seg[2:-2]):
                if not inner:
                    continue
                if inner.startswith("`") and inner.endswith("`") and len(inner) > 2:
                    r = par.add_run(inner[1:-1])
                    set_font(r, MONO, MONO, size - 1, bold=True, mono=True)
                else:
                    r = par.add_run(inner)
                    set_font(r, ascii_font, ea, size, bold=True)
        elif seg.startswith("`") and seg.endswith("`") and len(seg) > 2:
            r = par.add_run(seg[1:-1])
            set_font(r, MONO, MONO, size - 1, mono=True)
        else:
            r = par.add_run(seg)
            set_font(r, ascii_font, ea, size, bold=bold)


def field(par, instr, placeholder=""):
    r = par.add_run()

    def fld(t):
        el = OxmlElement("w:fldChar")
        el.set(qn("w:fldCharType"), t)
        return el

    r._element.append(fld("begin"))
    it = OxmlElement("w:instrText")
    it.set(qn("xml:space"), "preserve")
    it.text = instr
    r._element.append(it)
    r._element.append(fld("separate"))
    t = OxmlElement("w:t")
    t.text = placeholder
    r._element.append(t)
    r._element.append(fld("end"))
    return r


def shape_styles(doc):
    normal = doc.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal.font.size = Pt(12)
    normal.element.rPr.rFonts.set(qn("w:eastAsia"), SONG)
    normal.paragraph_format.line_spacing = 1.5
    normal.paragraph_format.space_after = Pt(0)

    sizes = {1: 16, 2: 14, 3: 13, 4: 12}
    for lvl, size in sizes.items():
        st = doc.styles[f"Heading {lvl}"]
        st.font.name = "Times New Roman"
        st.font.size = Pt(size)
        st.font.bold = True
        st.font.color.rgb = RGBColor(0, 0, 0)
        st.element.rPr.rFonts.set(qn("w:eastAsia"), HEI)
        st.paragraph_format.space_before = Pt(14 if lvl == 1 else 10)
        st.paragraph_format.space_after = Pt(8 if lvl == 1 else 6)
        st.paragraph_format.line_spacing = 1.4

    for name in ("List Bullet", "List Number"):
        st = doc.styles[name]
        st.font.name = "Times New Roman"
        st.font.size = Pt(12)
        st.element.rPr.rFonts.set(qn("w:eastAsia"), SONG)
        st.paragraph_format.line_spacing = 1.5


def page_setup(section):
    section.page_width, section.page_height = Cm(21), Cm(29.7)
    section.top_margin, section.bottom_margin = Cm(2.5), Cm(2.5)
    section.left_margin, section.right_margin = Cm(3.0), Cm(2.5)
    section.header_distance, section.footer_distance = Cm(1.5), Cm(1.5)


def cover(doc, doc_kind, date):
    for _ in range(3):
        doc.add_paragraph()
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_font(p.add_run(COVER["course"]), "Times New Roman", HEI, 22, bold=True)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(36)
    set_font(p.add_run(COVER["project"]), "Times New Roman", HEI, 26, bold=True)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(10)
    set_font(p.add_run(doc_kind), "Times New Roman", HEI, 20, bold=True)

    for _ in range(4):
        doc.add_paragraph()

    info = [("课程名称", COVER["course"]), ("项目名称", COVER["project"]),
            ("文档名称", doc_kind), ("项目成员", COVER["members"]),
            ("学　　号", COVER["sid"]), ("指导教师", COVER["teacher"]),
            ("完成日期", date)]
    t = doc.add_table(rows=0, cols=2)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False          # 固定布局，避免列宽被内容撑开
    for k, v in info:
        row = t.add_row().cells
        pk = row[0].paragraphs[0]
        pk.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        set_font(pk.add_run(k + "："), "Times New Roman", HEI, 14, bold=True)
        pv = row[1].paragraphs[0]
        set_font(pv.add_run(v), "Times New Roman", SONG, 14)
    t.columns[0].width = Cm(4.0)
    t.columns[1].width = Cm(11.5)


def toc_page(doc):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_font(p.add_run("目　　录"), "Times New Roman", HEI, 18, bold=True)
    doc.add_paragraph()
    p = doc.add_paragraph()
    field(p, 'TOC \\o "1-3" \\h \\z \\u',
          "（右键此处 →「更新域」→「更新整个目录」即可生成目录）")


def content_header_footer(section, doc_title):
    section.header.is_linked_to_previous = False
    section.footer.is_linked_to_previous = False

    hp = section.header.paragraphs[0]
    hp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_font(hp.add_run(f"{COVER['course']}　|　{doc_title}"),
             "Times New Roman", SONG, 9)
    pbdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "6")
    bottom.set(qn("w:color"), "999999")
    pbdr.append(bottom)
    hp._p.get_or_add_pPr().append(pbdr)

    fp = section.footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_font(fp.add_run("第 "), "Times New Roman", SONG, 10)
    field(fp, "PAGE", "1")
    set_font(fp.add_run(" 页"), "Times New Roman", SONG, 10)


# ============================================================ 表格列宽
def _eff_len(s):
    """按显示宽度估算字符串长度：CJK 记 2，其余记 1。"""
    return sum(2 if ord(ch) > 0x2E80 else 1 for ch in s)


def column_widths(header, rows, total_cm=15.5, min_cm=1.5, cap=48.0):
    """按各列内容的显示宽度分配列宽（长单元格按 cap 截断，防止一列吃掉整张表）。

    仅用于新文档（fit_tables=True）；默认路径不调用，保证既有产物字节不变。
    """
    ncol = len(header)
    scores = []
    for j in range(ncol):
        cells = [header[j] if j < len(header) else ""]
        cells += [r[j] if j < len(r) else "" for r in rows]
        scores.append(max(2.0, min(cap, max(_eff_len(c) for c in cells))))
    total = sum(scores) or 1.0
    widths = [max(min_cm, total_cm * x / total) for x in scores]
    for _ in range(8):
        over = sum(widths) - total_cm
        if over <= 0.01:
            break
        flex = [i for i, w in enumerate(widths) if w > min_cm + 0.001]
        pool = sum(widths[i] - min_cm for i in flex)
        if pool <= 0:
            break
        for i in flex:
            widths[i] -= over * (widths[i] - min_cm) / pool
    return widths


def fix_table(t, header, rows, min_cm=1.5):
    """固定表格布局：按内容分配列宽，并禁止行在分页处被切开。"""
    widths = column_widths(header, rows, min_cm=min_cm)
    t.autofit = False
    tblPr = t._tbl.tblPr
    layout = OxmlElement("w:tblLayout")
    layout.set(qn("w:type"), "fixed")
    tblPr.append(layout)
    for j, w in enumerate(widths):
        t.columns[j].width = Cm(w)
        for row in t.rows:
            row.cells[j].width = Cm(w)
    for row in t.rows:
        trPr = row._tr.get_or_add_trPr()
        cant = OxmlElement("w:cantSplit")
        trPr.insert(0, cant)
    return widths


# ============================================================ 正文构建
def build(md_name, out_name, doc_kind, anchors, date=None, fit_tables=False):
    blocks = parse_md(os.path.join(DOCS, md_name))
    doc = Document()
    shape_styles(doc)
    page_setup(doc.sections[0])
    cover(doc, doc_kind, date or DATES.get(out_name, ""))
    doc.add_page_break()
    toc_page(doc)

    sec = doc.add_section(WD_SECTION.NEW_PAGE)
    page_setup(sec)
    content_header_footer(sec, doc_kind)
    doc.sections[0].footer.is_linked_to_previous = False

    pending = list(anchors)
    used = set()
    fig_no = 0
    tbl_no = 0
    head_stack = {}

    def emit_figure(fname, caption):
        nonlocal fig_no
        fig_no += 1
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(8)
        p.paragraph_format.space_after = Pt(2)
        p.paragraph_format.keep_with_next = True
        p.add_run().add_picture(os.path.join(FIGS, fname), width=Cm(15))
        c = doc.add_paragraph()
        c.alignment = WD_ALIGN_PARAGRAPH.CENTER
        c.paragraph_format.space_after = Pt(12)
        set_font(c.add_run(f"图 {fig_no}　{caption}"), "Times New Roman", SONG, 10.5)

    for blk in blocks:
        # 图锚点：命中则在当前标题之前插入
        if blk[0] == "h":
            text = blk[2]
            for a in list(pending):
                if a[0] == text:
                    emit_figure(a[1], a[2])
                    used.add(a[0])
                    pending.remove(a)
            head_stack[blk[1]] = text
            for k in [k for k in head_stack if k > blk[1]]:
                del head_stack[k]
            doc.add_heading(text, level=min(blk[1], 4))
            continue

        if blk[0] == "p":
            p = doc.add_paragraph()
            p.paragraph_format.first_line_indent = Cm(0.74)
            p.paragraph_format.space_after = Pt(4)
            add_inline(p, blk[1])

        elif blk[0] == "quote":
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Cm(0.8)
            p.paragraph_format.space_after = Pt(6)
            add_inline(p, blk[1], size=10.5)

        elif blk[0] == "ul":
            for it in blk[1]:
                p = doc.add_paragraph(style="List Bullet")
                p.paragraph_format.line_spacing = 1.4
                add_inline(p, it)

        elif blk[0] == "ol":
            for it in blk[1]:
                p = doc.add_paragraph(style="List Number")
                p.paragraph_format.line_spacing = 1.4
                add_inline(p, it)

        elif blk[0] == "code":
            for line in blk[1].split("\n"):
                p = doc.add_paragraph()
                p.paragraph_format.line_spacing = 1.0
                p.paragraph_format.left_indent = Cm(0.6)
                p.paragraph_format.space_after = Pt(0)
                set_font(p.add_run(line if line else " "), MONO, MONO, 9, mono=True)

        elif blk[0] == "hr":
            continue

        elif blk[0] == "table":
            header, rows = blk[1], blk[2]
            ncol = len(header)
            cur_head = head_stack.get(max(head_stack) if head_stack else 0, "")
            if cur_head and max(head_stack) >= 2:
                tbl_no += 1
                c = doc.add_paragraph()
                c.alignment = WD_ALIGN_PARAGRAPH.CENTER
                c.paragraph_format.space_before = Pt(6)
                c.paragraph_format.space_after = Pt(3)
                c.paragraph_format.keep_with_next = True
                set_font(c.add_run(f"表 {tbl_no}　{cur_head}"),
                         "Times New Roman", SONG, 10.5)
            t = doc.add_table(rows=1, cols=ncol)
            t.style = "Table Grid"
            t.alignment = WD_TABLE_ALIGNMENT.CENTER
            for j, h in enumerate(header):
                cell = t.rows[0].cells[j]
                cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
                cell.paragraphs[0].paragraph_format.line_spacing = 1.15
                add_inline(cell.paragraphs[0], h, size=10.5, ea=HEI, bold=True)
            tr = t.rows[0]._tr
            trPr = tr.get_or_add_trPr()
            th = OxmlElement("w:tblHeader")
            trPr.append(th)
            for r in rows:
                cells = t.add_row().cells
                for j in range(ncol):
                    val = r[j] if j < len(r) else ""
                    cp = cells[j].paragraphs[0]
                    cp.paragraph_format.line_spacing = 1.15
                    add_inline(cp, val, size=10.5)
            if fit_tables:
                fix_table(t, header, rows)
            doc.add_paragraph().paragraph_format.space_after = Pt(2)

    missing = [a[0] for a in pending]
    if missing:
        raise SystemExit(f"[{md_name}] 找不到图锚点，请检查 md 标题是否变更：{missing}")

    out = os.path.join(OUTDIR, out_name)
    doc.save(out)
    print(f"OK {out_name}  图 {fig_no} 张 / 表 {tbl_no} 个 / {len(blocks)} 个块")
    return out


# ============================================================ 图锚点
REQ_ANCHORS = [
    ("3. 典型使用场景", "fig-usecase.png", "系统用例图"),
    ("4.1 老师端（FR-T）", "fig-function-tree.png", "系统功能结构图"),
    ("4.1 老师端（FR-T）", "fig-teacher-flow.png", "老师出卷业务流程"),
    ("4.1 老师端（FR-T）", "fig-student-flow.png", "学生答题业务流程"),
]

DESIGN_ANCHORS = [
    ("3.2 目录结构", "fig-architecture.png", "系统总体架构"),
    ("4.3 被否决的替代方案", "fig-state.png", "试卷状态机图"),
    ("5.2 表结构", "fig-class.png", "领域模型类图"),
    ("6. 出题引擎设计", "fig-er.png", "数据库 ER 图"),
    ("8. 前端设计", "fig-sequence.png", "学生提交作答与判分时序图"),
    ("11. 性能与并发", "fig-deployment.png", "系统部署图"),
]

TEST_ANCHORS = [
    ("2. 测试计划", "fig-test-levels.png", "测试分层与各层规模"),
    ("4. 测试用例", "fig-test-flow.png", "一次判据闭环：判据 → 失败 → 定位 → 修复 → 回归"),
    ("6. 缺陷记录", "fig-defect-stats.png", "缺陷按发现阶段与性质的分布"),
]

# 三份课程文档的构建清单：选择器 / md 源 / 产物 / 封面文档名 / 图锚点 / 是否按内容排列表格
# fit=True 的文档表格多且宽（测试文档有 7 列用例表与缺陷表），需要固定布局 + 按内容分配列宽，
# 并禁止行被分页切开。01/02 保持 fit=False：走的仍是改动前那条代码路径（不调用 fix_table），
# 所以重跑 01/02 的排版与既有产物一致，工具升级不会悄悄改掉已交付文档的表格。
DOC_JOBS = [
    ("01", "01-需求分析.md", "01-需求分析文档.docx", "需求分析文档", REQ_ANCHORS, False),
    ("02", "02-技术设计.md", "02-技术设计文档.docx", "技术设计文档", DESIGN_ANCHORS, False),
    ("03", "08-测试文档.md", "03-测试文档.docx", "测试文档", TEST_ANCHORS, True),
]


def main(argv=None):
    """不给参数＝构建全部（保持原行为）；给了选择器就只构建选中的，便于只重跑一份。"""
    argv = list(sys.argv[1:] if argv is None else argv)
    todo = DOC_JOBS
    if argv:
        keys = {a.strip().zfill(2) for a in argv}
        known = {j[0] for j in DOC_JOBS}
        unknown = sorted(k for k in keys if k not in known)
        if unknown:
            print("用法：python tools/build_docs.py [01|02|03 ...]（不给参数＝构建全部）")
            print("未知选择器：" + ", ".join(unknown))
            raise SystemExit(2)
        todo = [j for j in DOC_JOBS if j[0] in keys]
    os.makedirs(OUTDIR, exist_ok=True)
    for _, md_name, out_name, doc_kind, anchors, fit in todo:
        build(md_name, out_name, doc_kind, anchors, fit_tables=fit)


if __name__ == "__main__":
    main()
